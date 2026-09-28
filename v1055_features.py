# -*- coding: utf-8 -*-
"""Taxo 10.5-r5 — compatibility for newer «Шлях» XLSX workbook metadata.

Some exports contain a non-standard/forward-compatible ``tabId`` attribute on
``<sheet>`` records in ``xl/workbook.xml``. openpyxl 3.1.5 rejects that
attribute while constructing ChildSheet. Taxo retries only this exact failure
using an in-memory copy with the unsupported attribute removed. The user's
source XLSX is never modified.
"""
from __future__ import annotations

import io
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

import vehicle_registry

APP_VERSION = "10.5-r5"


def _is_childsheet_tabid_error(exc):
    text = str(exc)
    return isinstance(exc, TypeError) and "unexpected keyword argument" in text and "tabId" in text


def _xlsx_without_sheet_tabid(path):
    """Return an in-memory XLSX copy without only the unsupported sheet tabId."""
    source = Path(path)
    out = io.BytesIO()
    changed = False
    with zipfile.ZipFile(str(source), "r") as src, zipfile.ZipFile(out, "w") as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "xl/workbook.xml":
                root = ET.fromstring(data)
                for elem in root.iter():
                    if elem.tag.rsplit("}", 1)[-1] != "sheet":
                        continue
                    for key in list(elem.attrib):
                        if key.rsplit("}", 1)[-1] == "tabId":
                            del elem.attrib[key]
                            changed = True
                if changed:
                    data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            dst.writestr(info, data)
    if not changed:
        out.close()
        raise ValueError("XLSX містить помилку ChildSheet/tabId, але атрибут tabId у workbook.xml не знайдено.")
    out.seek(0)
    return out


def parse_registry_xlsx_compat(path):
    """Parse «Шлях» XLSX and narrowly recover from the ChildSheet/tabId defect."""
    from openpyxl import load_workbook

    path = Path(path)
    compat_stream = None
    wb = None
    try:
        try:
            wb = load_workbook(filename=str(path), read_only=True, data_only=True)
        except TypeError as exc:
            if not _is_childsheet_tabid_error(exc):
                raise
            compat_stream = _xlsx_without_sheet_tabid(path)
            wb = load_workbook(filename=compat_stream, read_only=True, data_only=True)

        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            header_row, mapping = vehicle_registry._detect_header(rows)
            if header_row is None:
                continue
            parsed = []
            for row_no, row in enumerate(rows[header_row + 1:], start=header_row + 2):
                if not any(vehicle_registry._text(v) for v in row):
                    continue
                item = vehicle_registry._record(row, mapping, row_no)
                if not item["plate"] and not item["vin"]:
                    continue
                parsed.append(item)
            return {
                "source_kind": vehicle_registry.SOURCE_KIND,
                "source_name": path.name,
                "sheet_name": ws.title,
                "header_row": header_row + 1,
                "rows": parsed,
            }
    finally:
        if wb is not None:
            wb.close()
        if compat_stream is not None:
            compat_stream.close()
    raise ValueError("Формат XLSX не розпізнано як експорт транспортних засобів з реєстру «Шлях».")


def install(core, base_app):
    if getattr(core, "_TAXO_1055_INSTALLED", False):
        return core.App

    vehicle_registry.parse_registry_xlsx = parse_registry_xlsx_compat
    core.APP_VERSION = APP_VERSION

    class Taxo1055App(base_app):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            core.APP_VERSION = APP_VERSION
            self.title("Taxo %s — Працівники, графіки та шляхівки" % core.APP_VERSION)

    Taxo1055App.__name__ = "App"
    Taxo1055App.__qualname__ = "App"
    core.App = Taxo1055App
    core._TAXO_1055_INSTALLED = True
    return Taxo1055App
