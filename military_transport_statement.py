# -*- coding: utf-8 -*-
"""Taxo 10.5-r1 — official Appendix 1 military-transport statement.

The statement is an employer-side report for the enterprise's own/balance fleet.
It is intentionally independent from:
- the registry «Шлях» status;
- operational vehicle document control;
- military transport orders / transfer status;
- operational ``vehicles.active``.

Unknown ownership/report scope is never guessed. A vehicle enters the official
statement only after an explicit ``own`` scope decision.
"""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import military_transport_2026 as mt
import personnel_registry as personnel_registry

APP_VERSION = "10.5-r1"
LEGAL_BASIS = mt.LEGAL_BASIS
FORM_NAME = (
    "Відомість про наявність і технічний стан транспортних засобів і техніки, "
    "а також про громадян, які працюють на підприємстві, в установі та організації "
    "на таких транспортних засобах і техніці"
)
FORM_REFERENCE = "Додаток 1 до Положення про військово-транспортний обов'язок"

SCOPE_UNKNOWN = "unknown"
SCOPE_OWN = "own"
SCOPE_EXCLUDED = "excluded"
SCOPE_LABELS = {
    SCOPE_UNKNOWN: "Потрібно визначити",
    SCOPE_OWN: "Власний / балансовий парк",
    SCOPE_EXCLUDED: "Не включати до власного парку",
}
VALID_SCOPES = frozenset(SCOPE_LABELS)

HEADERS = (
    "№ з/п",
    "Тип",
    "Марка",
    "Державний номер",
    "Технічний стан",
    "Рік випуску",
    "Залишкова (балансова) вартість, тис. грн",
    "Прізвище, ім'я, по батькові працівника",
    "Рік народження",
    "Номер військово-облікової спеціальності",
    "Військове звання",
    "Місце проживання",
)


def _iso_day(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value or "").strip()
    if not text:
        return date.today().isoformat()
    return date.fromisoformat(text[:10]).isoformat()


def _table_exists(con, table):
    return con.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name=?",
        (table,),
    ).fetchone() is not None


def _columns(con, table):
    if not _table_exists(con, table):
        return set()
    return {row[1] for row in con.execute("PRAGMA table_info(%s)" % table).fetchall()}


def ensure_schema_on_connection(con):
    mt.ensure_schema_on_connection(con)
    personnel_registry.ensure_schema_on_connection(con)
    con.executescript(
        """
        CREATE TABLE IF NOT EXISTS vehicle_military_transport_statement_data (
            vehicle_id INTEGER PRIMARY KEY REFERENCES vehicles(id) ON DELETE CASCADE,
            report_scope TEXT NOT NULL DEFAULT 'unknown',
            vehicle_type TEXT DEFAULT '',
            technical_condition TEXT DEFAULT '',
            residual_book_value_thousand_uah TEXT DEFAULT '',
            note TEXT DEFAULT '',
            updated_at TEXT NOT NULL
        );
        CREATE INDEX IF NOT EXISTS idx_vehicle_military_statement_scope
            ON vehicle_military_transport_statement_data(report_scope,vehicle_id);
        """
    )


def ensure_schema(core):
    con = core.db()
    try:
        ensure_schema_on_connection(con)
        con.commit()
    finally:
        con.close()


def get_vehicle_statement_data(con, vehicle_id):
    ensure_schema_on_connection(con)
    return con.execute(
        "SELECT * FROM vehicle_military_transport_statement_data WHERE vehicle_id=?",
        (int(vehicle_id),),
    ).fetchone()


def set_vehicle_statement_data(
    con,
    vehicle_id,
    *,
    report_scope=SCOPE_UNKNOWN,
    vehicle_type="",
    technical_condition="",
    residual_book_value_thousand_uah="",
    note="",
):
    ensure_schema_on_connection(con)
    if report_scope not in VALID_SCOPES:
        raise ValueError("Невідомий статус включення ТЗ до відомості.")
    value = str(residual_book_value_thousand_uah or "").strip().replace(",", ".")
    if value:
        try:
            numeric = float(value)
        except ValueError as exc:
            raise ValueError("Залишкова балансова вартість має бути числом у тис. грн.") from exc
        if numeric < 0:
            raise ValueError("Залишкова балансова вартість не може бути від'ємною.")
        value = ("%.3f" % numeric).rstrip("0").rstrip(".")
    now = datetime.now().isoformat(timespec="seconds")
    con.execute(
        """INSERT INTO vehicle_military_transport_statement_data(
               vehicle_id,report_scope,vehicle_type,technical_condition,
               residual_book_value_thousand_uah,note,updated_at
           ) VALUES(?,?,?,?,?,?,?)
           ON CONFLICT(vehicle_id) DO UPDATE SET
               report_scope=excluded.report_scope,
               vehicle_type=excluded.vehicle_type,
               technical_condition=excluded.technical_condition,
               residual_book_value_thousand_uah=excluded.residual_book_value_thousand_uah,
               note=excluded.note,
               updated_at=excluded.updated_at""",
        (
            int(vehicle_id),
            report_scope,
            str(vehicle_type or "").strip(),
            str(technical_condition or "").strip(),
            value,
            str(note or "").strip(),
            now,
        ),
    )


def _vehicle_local_value(row, columns, *names):
    for name in names:
        if name in columns:
            value = row[name]
            if value not in (None, ""):
                return str(value).strip()
    return ""


def _vehicle_base_rows(con):
    ensure_schema_on_connection(con)
    vcols = _columns(con, "vehicles")
    rows = con.execute(
        """SELECT v.*,s.report_scope,s.vehicle_type AS statement_vehicle_type,
                  s.technical_condition AS statement_technical_condition,
                  s.residual_book_value_thousand_uah,s.note AS statement_note,
                  p.readiness_status
           FROM vehicles v
           LEFT JOIN vehicle_military_transport_statement_data s ON s.vehicle_id=v.id
           LEFT JOIN vehicle_military_transport_profile p ON p.vehicle_id=v.id
           ORDER BY COALESCE(v.plate,''),v.name,v.id"""
    ).fetchall()
    return rows, vcols


def _vehicle_identity(row, vcols):
    vehicle_type = str(row["statement_vehicle_type"] or "").strip()
    if not vehicle_type:
        vehicle_type = _vehicle_local_value(row, vcols, "registry_vehicle_type")
    make = ""
    local_make = _vehicle_local_value(row, vcols, "make")
    local_model = _vehicle_local_value(row, vcols, "model")
    if local_make or local_model:
        make = " ".join(x for x in (local_make, local_model) if x)
    if not make:
        make = _vehicle_local_value(row, vcols, "make_model", "name")
    technical = str(row["statement_technical_condition"] or "").strip()
    if not technical:
        readiness = str(row["readiness_status"] or mt.READINESS_UNKNOWN)
        if readiness != mt.READINESS_UNKNOWN:
            technical = mt.READINESS_LABELS.get(readiness, readiness)
    return {
        "vehicle_type": vehicle_type,
        "make": make,
        "plate": str(row["plate"] or "").strip(),
        "year": str(row["year"] or "").strip(),
        "technical_condition": technical,
        "residual_value": str(row["residual_book_value_thousand_uah"] or "").strip(),
    }


def statement_overview(con, as_of=None):
    """Return every Taxo vehicle with explicit report-scope state.

    ``vehicles.active`` is deliberately ignored: operational activity is not an
    ownership/balance criterion.
    """
    day = _iso_day(as_of or date.today())
    rows, vcols = _vehicle_base_rows(con)
    result = []
    for row in rows:
        scope = str(row["report_scope"] or SCOPE_UNKNOWN)
        if scope not in VALID_SCOPES:
            scope = SCOPE_UNKNOWN
        identity = _vehicle_identity(row, vcols)
        worker_count = len(_workers_for_statement(con, row["id"], day))
        result.append({
            "vehicle_id": int(row["id"]),
            "name": str(row["name"] or ""),
            "plate": identity["plate"],
            "report_scope": scope,
            "scope_label": SCOPE_LABELS[scope],
            "vehicle_type": identity["vehicle_type"],
            "technical_condition": identity["technical_condition"],
            "residual_value": identity["residual_value"],
            "worker_count": worker_count,
            "statement_note": str(row["statement_note"] or ""),
        })
    return result


def _workers_for_statement(con, vehicle_id, as_of):
    day = _iso_day(as_of)
    return con.execute(
        """SELECT w.*,e.last_name,e.first_name,e.middle_name,e.birth_date,
                  e.actual_address,e.registered_address,
                  m.military_specialty,m.military_rank
           FROM vehicle_military_transport_workers w
           LEFT JOIN employees e ON e.id=w.employee_id
           LEFT JOIN employee_military_profile m ON m.employee_id=w.employee_id
           WHERE w.vehicle_id=?
             AND (COALESCE(w.valid_from,'')='' OR w.valid_from<=?)
             AND (COALESCE(w.valid_until,'')='' OR w.valid_until>=?)
           ORDER BY COALESCE(e.last_name,w.worker_name),e.first_name,e.middle_name,w.id""",
        (int(vehicle_id), day, day),
    ).fetchall()


def collect_statement_rows(con, as_of=None):
    """Collect official Appendix 1 rows for explicitly marked own/balance vehicles."""
    day = _iso_day(as_of or date.today())
    rows, vcols = _vehicle_base_rows(con)
    output = []
    seq = 0
    for vehicle in rows:
        scope = str(vehicle["report_scope"] or SCOPE_UNKNOWN)
        if scope != SCOPE_OWN:
            continue
        seq += 1
        identity = _vehicle_identity(vehicle, vcols)
        workers = _workers_for_statement(con, vehicle["id"], day)
        if not workers:
            workers = [None]
        for index, worker in enumerate(workers):
            if worker is None:
                worker_name = birth_year = specialty = rank = residence = ""
            else:
                worker_name = str(worker["worker_name"] or "").strip()
                if not worker_name:
                    worker_name = " ".join(
                        x for x in (
                            str(worker["last_name"] or "").strip(),
                            str(worker["first_name"] or "").strip(),
                            str(worker["middle_name"] or "").strip(),
                        ) if x
                    )
                birth = str(worker["birth_date"] or "").strip()
                birth_year = birth[:4] if len(birth) >= 4 else ""
                specialty = str(worker["military_specialty"] or "").strip()
                rank = str(worker["military_rank"] or "").strip()
                residence = str(worker["actual_address"] or worker["registered_address"] or "").strip()
            output.append({
                "seq": seq if index == 0 else "",
                "vehicle_id": int(vehicle["id"]),
                "vehicle_type": identity["vehicle_type"],
                "make": identity["make"],
                "plate": identity["plate"],
                "technical_condition": identity["technical_condition"],
                "year": identity["year"],
                "residual_value": identity["residual_value"],
                "worker_name": worker_name,
                "birth_year": birth_year,
                "military_specialty": specialty,
                "military_rank": rank,
                "residence": residence,
            })
    return output


def statement_issues(con, as_of=None):
    day = _iso_day(as_of or date.today())
    overview = statement_overview(con, day)
    rows = collect_statement_rows(con, day)
    issues = []
    unknown = [row for row in overview if row["report_scope"] == SCOPE_UNKNOWN]
    if unknown:
        issues.append(
            "Не визначено належність до власного/балансового парку для %d ТЗ." % len(unknown)
        )
    own = [row for row in overview if row["report_scope"] == SCOPE_OWN]
    if not own:
        issues.append("Не позначено жодного ТЗ як власний / балансовий парк.")
    for row in own:
        missing = []
        if not row["vehicle_type"]:
            missing.append("тип")
        if not row["technical_condition"]:
            missing.append("технічний стан")
        if not row["residual_value"]:
            missing.append("залишкова вартість")
        if row["worker_count"] == 0:
            missing.append("працівник")
        if missing:
            issues.append(
                "%s: не заповнено %s." % (row["plate"] or row["name"], ", ".join(missing))
            )
    for row in rows:
        if row["worker_name"] and not row["birth_year"]:
            issues.append("%s: для працівника не вказано рік народження." % (row["plate"] or "ТЗ"))
        if row["worker_name"] and not row["military_specialty"]:
            issues.append("%s: для працівника не вказано ВОС." % (row["plate"] or "ТЗ"))
        if row["worker_name"] and not row["military_rank"]:
            issues.append("%s: для працівника не вказано військове звання." % (row["plate"] or "ТЗ"))
        if row["worker_name"] and not row["residence"]:
            issues.append("%s: для працівника не вказано місце проживання." % (row["plate"] or "ТЗ"))
    return list(dict.fromkeys(issues))


def company_details(con):
    if not _table_exists(con, "company"):
        return {}
    row = con.execute("SELECT * FROM company WHERE id=1").fetchone()
    if row is None:
        return {}
    return {key: row[key] for key in row.keys()}


def _company_line(company):
    parts = [
        str(company.get("name") or "").strip(),
        str(company.get("address") or "").strip(),
    ]
    contacts = ", ".join(
        value for value in (
            str(company.get("phone") or "").strip(),
            str(company.get("email") or "").strip(),
        ) if value
    )
    if contacts:
        parts.append(contacts)
    return " · ".join(part for part in parts if part)


def export_statement_xlsx(path, report_date, company, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Відомість"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    ws.cell(1, 1, FORM_NAME)
    ws.cell(1, 1).font = Font(bold=True, size=12)
    ws.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 42

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(HEADERS))
    ws.cell(2, 1, "%s · станом на %s" % (FORM_REFERENCE, date.fromisoformat(_iso_day(report_date)).strftime("%d.%m.%Y")))
    ws.cell(2, 1).alignment = Alignment(horizontal="center", wrap_text=True)

    company_line = _company_line(company)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(HEADERS))
    ws.cell(3, 1, company_line)
    ws.cell(3, 1).alignment = Alignment(horizontal="left", wrap_text=True)

    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(5, col, header)
        cell.font = Font(bold=True, size=9)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[5].height = 58

    for row_no, item in enumerate(rows, start=6):
        values = (
            item["seq"], item["vehicle_type"], item["make"], item["plate"],
            item["technical_condition"], item["year"], item["residual_value"],
            item["worker_name"], item["birth_year"], item["military_specialty"],
            item["military_rank"], item["residence"],
        )
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row_no, col, value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    widths = (7, 15, 22, 15, 18, 11, 18, 28, 12, 18, 16, 30)
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    signature_row = 7 + len(rows)
    signer_position = str(company.get("signer_position") or "").strip() or "Керівник"
    signer_name = str(company.get("signer_name") or "").strip()
    ws.merge_cells(start_row=signature_row, start_column=1, end_row=signature_row, end_column=5)
    ws.cell(signature_row, 1, signer_position)
    ws.merge_cells(start_row=signature_row, start_column=8, end_row=signature_row, end_column=12)
    ws.cell(signature_row, 8, signer_name)

    ws.freeze_panes = "A6"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:5"
    ws.print_area = "A1:L%d" % signature_row
    ws.oddFooter.center.text = "Taxo · %s" % LEGAL_BASIS
    wb.save(path)
    return path


def _report_font_path(preferred=None):
    candidates = []
    if preferred:
        candidates.append(str(preferred))
    candidates.extend([
        r"C:\Windows\Fonts\arial.ttf",
        r"C:\Windows\Fonts\calibri.ttf",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Times New Roman.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ])
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            return candidate
    return ""


def export_statement_pdf(path, report_date, company, rows, font_path=None):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    font_name = "Helvetica"
    resolved_font = _report_font_path(font_path)
    if resolved_font:
        font_name = "TaxoStatement"
        pdfmetrics.registerFont(TTFont(font_name, resolved_font))

    doc = SimpleDocTemplate(
        str(path),
        pagesize=landscape(A4),
        leftMargin=7 * mm,
        rightMargin=7 * mm,
        topMargin=7 * mm,
        bottomMargin=8 * mm,
    )
    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "TaxoStatementTitle", parent=styles["Normal"], fontName=font_name,
        fontSize=9, leading=11, alignment=TA_CENTER, spaceAfter=4,
    )
    small = ParagraphStyle(
        "TaxoStatementSmall", parent=styles["Normal"], fontName=font_name,
        fontSize=6.5, leading=8,
    )
    header_style = ParagraphStyle(
        "TaxoStatementHeader", parent=small, alignment=TA_CENTER,
    )

    story = [
        Paragraph(FORM_NAME, title),
        Paragraph(
            "%s · станом на %s" % (
                FORM_REFERENCE,
                date.fromisoformat(_iso_day(report_date)).strftime("%d.%m.%Y"),
            ),
            title,
        ),
    ]
    company_line = _company_line(company)
    if company_line:
        story.append(Paragraph(company_line, small))
    story.append(Spacer(1, 3 * mm))

    data = [[Paragraph(h, header_style) for h in HEADERS]]
    for item in rows:
        values = (
            item["seq"], item["vehicle_type"], item["make"], item["plate"],
            item["technical_condition"], item["year"], item["residual_value"],
            item["worker_name"], item["birth_year"], item["military_specialty"],
            item["military_rank"], item["residence"],
        )
        data.append([Paragraph(str(value or ""), small) for value in values])

    widths = [8, 18, 25, 18, 22, 12, 25, 35, 15, 24, 22, 36]
    total = sum(widths)
    available = landscape(A4)[0] - 14 * mm
    col_widths = [available * w / total for w in widths]
    table = Table(data, colWidths=col_widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), font_name),
        ("FONTSIZE", (0, 0), (-1, -1), 6.5),
        ("LEADING", (0, 0), (-1, -1), 8),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.black),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
    ]))
    story.append(table)
    story.append(Spacer(1, 4 * mm))
    signer_position = str(company.get("signer_position") or "").strip() or "Керівник"
    signer_name = str(company.get("signer_name") or "").strip()
    story.append(Paragraph("%s ____________________ %s" % (signer_position, signer_name), small))
    doc.build(story)
    return path
