# -*- coding: utf-8 -*-
"""Taxo 10.5-r4 — lossless state-registry snapshots for employees.

Canonical Taxo fields remain validated working data.  This layer additionally
preserves the exact value of every source XLSX column so that a value rejected by
canonical validation is still visible as "what the state registry currently says".
Missing/blank source values never clear canonical Taxo data.
"""
from __future__ import annotations

from datetime import datetime

APP_VERSION = "10.5-r4"

STATUS_MAPPED = "mapped"
STATUS_REJECTED = "rejected_format"
STATUS_IDENTITY = "identity"
STATUS_BLANK = "blank"
STATUS_SOURCE_ONLY = "source_only"

STATUS_LABELS = {
    STATUS_MAPPED: "Перенесено / доступно для звірки",
    STATUS_REJECTED: "Не перенесено через формат",
    STATUS_IDENTITY: "Ідентифікаційне поле",
    STATUS_BLANK: "Порожньо у витягу",
    STATUS_SOURCE_ONLY: "Лише у витягу",
}


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _norm_header(registry, value):
    return registry._norm(value)  # reuse the parser's apostrophe/spacing normalization


def _field_maps(registry):
    detailed = {
        "прізвище імя по батькові": ("identity", "full_name"),
        "рнокпп": ("personal", "rnokpp"),
        "серія та номер паспорту": ("personal", "passport_number"),
        "номер id картки": ("personal", "id_card_number"),
        "дата народження": ("personal", "birth_date"),
        "адреса зареєстрованого місця проживання": ("personal", "registered_address"),
        "адреса фактичного місця проживання": ("personal", "actual_address"),
        "ідентифікатор військовозобовязаного": ("military", "registry_person_id"),
        "вид обліку особи": ("military", "account_type"),
        "статус обліку особи": ("military", "account_status"),
        "належність до резервістів": ("military", "reservist_status"),
        "бронювання": ("military", "booking_status"),
        "дата кінця бронювання": ("military", "booking_until"),
        "причина відстрочки": ("military", "deferment_reason"),
        "дата закінчення відстрочки": ("military", "deferment_until"),
        "військове звання": ("military", "military_rank"),
        "військово облікова спеціальність": ("military", "military_specialty"),
        "найменування тцк органу сбу відповідного підрозділу розвідувального органу в якому перебуває на військовому обліку": ("military", "account_authority"),
        "військова служба": ("military", "military_service_status"),
    }
    summary = {
        "прізвище": ("identity", "last_name"),
        "імя": ("identity", "first_name"),
        "по батькові": ("identity", "middle_name"),
        "стать": ("personal", "gender"),
        "дата народження": ("personal", "birth_date"),
        "рнокпп": ("personal", "rnokpp"),
        "серія паспорту": ("personal", "passport_series"),
        "номер паспорту": ("personal", "passport_number"),
        "номер id картки": ("personal", "id_card_number"),
        "статус": ("military", "booking_status"),
        "примітка": ("military", "registry_note"),
    }
    return {
        registry.SOURCE_DETAILED: detailed,
        registry.SOURCE_SUMMARY: summary,
    }


def ensure_schema_on_connection(registry, original_ensure, con):
    original_ensure(con)
    cols = {row[1] for row in con.execute("PRAGMA table_info(employee_military_profile)").fetchall()}
    if "registry_note" not in cols:
        con.execute("ALTER TABLE employee_military_profile ADD COLUMN registry_note TEXT DEFAULT ''")
    con.executescript("""
        CREATE TABLE IF NOT EXISTS employee_registry_raw_fields (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            import_id INTEGER NOT NULL REFERENCES employee_registry_imports(id) ON DELETE CASCADE,
            employee_id INTEGER REFERENCES employees(id) ON DELETE SET NULL,
            source_kind TEXT NOT NULL,
            source_name TEXT NOT NULL,
            file_sha256 TEXT NOT NULL,
            source_row INTEGER,
            field_order INTEGER NOT NULL,
            field_key TEXT NOT NULL,
            source_header TEXT NOT NULL,
            raw_value TEXT DEFAULT '',
            canonical_scope TEXT DEFAULT '',
            canonical_field TEXT DEFAULT '',
            canonical_value TEXT DEFAULT '',
            transfer_status TEXT NOT NULL,
            transfer_note TEXT DEFAULT '',
            imported_at TEXT NOT NULL,
            UNIQUE(import_id,source_row,field_order)
        );
        CREATE INDEX IF NOT EXISTS idx_employee_registry_raw_employee
            ON employee_registry_raw_fields(employee_id,import_id DESC,field_order);
        CREATE INDEX IF NOT EXISTS idx_employee_registry_raw_import
            ON employee_registry_raw_fields(import_id,source_row,field_order);
    """)


def _canonical_value(item, scope, field):
    if scope == "identity":
        if field == "full_name":
            return _text(item.get("full_name"))
        return _text(item.get(field))
    return _text((item.get(scope) or {}).get(field))


def _classify_field(item, scope, field, raw):
    if not raw:
        return STATUS_BLANK, ""
    if scope == "identity":
        return STATUS_IDENTITY, "Використовується для зіставлення особи; автоматично ПІБ не переписується."
    canonical = _canonical_value(item, scope, field)
    if canonical:
        return STATUS_MAPPED, ""
    return STATUS_REJECTED, "Оригінальне значення збережено, але canonical-поле не заповнено через перевірку формату."


def _attach_raw_fields(registry, parsed, path):
    from openpyxl import load_workbook

    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        ws = wb[parsed["sheet_name"]]
        rows = list(ws.iter_rows(values_only=True))
        header_index = int(parsed["header_row"]) - 1
        headers = rows[header_index] if 0 <= header_index < len(rows) else ()
        mappings = _field_maps(registry).get(parsed["source_kind"], {})
        by_source_row = {int(item.get("source_row")): item for item in parsed.get("rows") or [] if item.get("source_row")}
        for source_row, item in by_source_row.items():
            row = rows[source_row - 1] if 0 < source_row <= len(rows) else ()
            # The summary form's free-text note is a useful official value in its own right.
            if parsed["source_kind"] == registry.SOURCE_SUMMARY:
                note_idx = None
                for idx, header in enumerate(headers):
                    if _norm_header(registry, header) == "примітка":
                        note_idx = idx
                        break
                if note_idx is not None and note_idx < len(row):
                    note = _text(row[note_idx])
                    if note:
                        item.setdefault("military", {})["registry_note"] = note
            raw_fields = []
            for idx, header in enumerate(headers):
                header_text = _text(header)
                if not header_text:
                    continue
                raw = _text(row[idx]) if idx < len(row) else ""
                key = _norm_header(registry, header_text)
                mapping = mappings.get(key)
                if mapping is None:
                    scope = field = canonical = ""
                    status = STATUS_BLANK if not raw else STATUS_SOURCE_ONLY
                    note = "" if not raw else "Колонка є у державному витягу, але не має canonical-поля Taxo."
                else:
                    scope, field = mapping
                    canonical = _canonical_value(item, scope, field)
                    status, note = _classify_field(item, scope, field, raw)
                raw_fields.append({
                    "field_order": idx,
                    "field_key": key,
                    "source_header": header_text,
                    "raw_value": raw,
                    "canonical_scope": scope,
                    "canonical_field": field,
                    "canonical_value": canonical,
                    "transfer_status": status,
                    "transfer_note": note,
                })
            item["raw_fields"] = raw_fields
        return parsed
    finally:
        wb.close()


def latest_raw_snapshot(con, employee_id):
    row = con.execute(
        """SELECT import_id, MAX(imported_at) AS imported_at
             FROM employee_registry_raw_fields
            WHERE employee_id=?
         GROUP BY import_id
         ORDER BY imported_at DESC, import_id DESC LIMIT 1""",
        (int(employee_id),),
    ).fetchone()
    if row is None:
        return []
    return con.execute(
        """SELECT * FROM employee_registry_raw_fields
            WHERE employee_id=? AND import_id=?
            ORDER BY source_row,field_order,id""",
        (int(employee_id), int(row["import_id"])),
    ).fetchall()


def list_raw_snapshot_history(con, employee_id, limit=500):
    return con.execute(
        """SELECT * FROM employee_registry_raw_fields
            WHERE employee_id=?
            ORDER BY imported_at DESC,import_id DESC,source_row,field_order
            LIMIT ?""",
        (int(employee_id), int(limit)),
    ).fetchall()


def install(registry):
    """Patch the existing registry module without forking its working data model."""
    if getattr(registry, "_TAXO_1054_LOSSLESS", False):
        return registry

    original_ensure = registry.ensure_schema_on_connection
    original_parse = registry.parse_registry_xlsx
    original_apply = registry.apply_registry_import

    registry.MILITARY_FIELD_LABELS.setdefault("registry_note", "Примітка державного витягу")

    def ensure_schema(con):
        return ensure_schema_on_connection(registry, original_ensure, con)

    def parse_registry_xlsx(path):
        parsed = original_parse(path)
        return _attach_raw_fields(registry, parsed, path)

    def apply_registry_import(core, preview, mode=registry.IMPORT_UPDATE):
        result = original_apply(core, preview, mode=mode)
        con = core.db()
        try:
            ensure_schema(con)
            now = datetime.now().isoformat(timespec="seconds")
            parsed = preview.get("parsed") or {}
            for row_plan in preview.get("plan") or []:
                item = row_plan.get("item") or {}
                employee_id = row_plan.get("employee_id")
                for raw in item.get("raw_fields") or []:
                    con.execute(
                        """INSERT OR REPLACE INTO employee_registry_raw_fields(
                               import_id,employee_id,source_kind,source_name,file_sha256,source_row,
                               field_order,field_key,source_header,raw_value,canonical_scope,
                               canonical_field,canonical_value,transfer_status,transfer_note,imported_at
                           ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                        (result["import_id"], employee_id, parsed.get("source_kind", ""),
                         parsed.get("source_name", ""), preview.get("file_sha256", ""),
                         item.get("source_row"), int(raw.get("field_order", 0)), raw.get("field_key", ""),
                         raw.get("source_header", ""), raw.get("raw_value", ""),
                         raw.get("canonical_scope", ""), raw.get("canonical_field", ""),
                         raw.get("canonical_value", ""), raw.get("transfer_status", STATUS_SOURCE_ONLY),
                         raw.get("transfer_note", ""), now),
                    )
            con.commit()
        except Exception:
            con.rollback()
            raise
        finally:
            con.close()
        return result

    registry.ensure_schema_on_connection = ensure_schema
    registry.parse_registry_xlsx = parse_registry_xlsx
    registry.apply_registry_import = apply_registry_import
    registry._TAXO_1054_LOSSLESS = True
    return registry
