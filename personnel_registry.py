# -*- coding: utf-8 -*-
"""Taxo 10.4-r5 — personnel registry import, military profile and document register.

The module intentionally keeps personal XLSX files outside the repository. Runtime
imports store structured data and provenance in the user's working SQLite database.
No employee is silently created from a registry extract.
"""
from __future__ import annotations

import hashlib
import json
import re
from datetime import date, datetime
from pathlib import Path

APP_VERSION = "10.4-r5"
SOURCE_DETAILED = "military_registry_personal"
SOURCE_SUMMARY = "military_registry_employees"

IMPORT_COMPARE = "compare"
IMPORT_FILL_EMPTY = "fill_empty"
IMPORT_UPDATE = "update"
IMPORT_MODES = (IMPORT_COMPARE, IMPORT_FILL_EMPTY, IMPORT_UPDATE)
IMPORT_MODE_LABELS = {
    IMPORT_COMPARE: "Лише звірити",
    IMPORT_FILL_EMPTY: "Доповнити порожні поля",
    IMPORT_UPDATE: "Оновити з реєстру + доповнити",
}

DOCUMENT_TYPES = (
    "Паспорт громадянина України",
    "ID-картка",
    "Військово-обліковий документ",
    "Посвідчення водія",
    "РНОКПП",
    "Документ про освіту",
    "Медичний документ",
    "Інше",
)

EMPLOYEE_FIELD_LABELS = {
    "gender": "Стать",
    "birth_date": "Дата народження",
    "rnokpp": "РНОКПП",
    "passport_series": "Серія паспорта",
    "passport_number": "Номер паспорта",
    "id_card_number": "Номер ID-картки",
    "registered_address": "Зареєстроване місце проживання",
    "actual_address": "Фактичне місце проживання",
}

MILITARY_FIELD_LABELS = {
    "registry_person_id": "Ідентифікатор у Реєстрі",
    "account_type": "Вид обліку особи",
    "account_status": "Статус обліку особи",
    "reservist_status": "Належність до резервістів",
    "booking_status": "Бронювання",
    "booking_until": "Дата кінця бронювання",
    "deferment_reason": "Причина відстрочки",
    "deferment_until": "Дата закінчення відстрочки",
    "military_rank": "Військове звання",
    "military_specialty": "Військово-облікова спеціальність",
    "account_authority": "ТЦК / орган обліку",
    "military_service_status": "Військова служба",
    "reserve_category": "Категорія запасу / склад",
    "military_document": "Військово-обліковий документ",
    "education_specialty": "Освіта / спеціальність",
    "foreign_passport": "Паспорт для виїзду за кордон",
    "fitness": "Придатність до військової служби",
    "family_info": "Сімейний стан / члени сім'ї",
    "appointment_act": "Акт про призначення / звільнення",
    "notification_requisites": "Повідомлення про призначення / звільнення",
}

# Superset of the employer-side Appendix 5 data groups. Some fields depend on
# the person's category and may legitimately be inapplicable; UI therefore
# reports them as "not filled", not as an automatic legal violation.
APPENDIX5_PROFILE_FIELDS = (
    "registry_person_id",
    "military_rank",
    "military_specialty",
    "reserve_category",
    "military_document",
    "education_specialty",
    "foreign_passport",
    "account_authority",
    "reservist_status",
    "fitness",
    "family_info",
    "appointment_act",
    "notification_requisites",
)

_DETAILED_HEADERS = {
    "прізвище імя по батькові",
    "рнокпп",
    "дата народження",
    "ідентифікатор військовозобовязаного",
    "вид обліку особи",
    "статус обліку особи",
}
_SUMMARY_HEADERS = {
    "прізвище",
    "імя",
    "по батькові",
    "стать",
    "дата народження",
    "рнокпп",
    "статус",
    "примітка",
}


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _norm(value):
    text=_text(value).lower().replace("ʼ", "'").replace("’", "'").replace("`", "'")
    text=text.replace("'", "").replace("-", " ")
    text=re.sub(r"[^0-9a-zа-яіїєґ]+", " ", text, flags=re.IGNORECASE)
    return re.sub(r"\s+", " ", text).strip()


def _name_key(last_name, first_name, middle_name=""):
    return _norm(" ".join(x for x in (_text(last_name),_text(first_name),_text(middle_name)) if x))


def _split_full_name(value):
    parts=[x for x in re.split(r"\s+", _text(value)) if x]
    if not parts:
        return "", "", ""
    if len(parts)==1:
        return parts[0], "", ""
    return parts[0], parts[1], " ".join(parts[2:])


def _iso_date(value):
    if value in (None, ""):
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text=_text(value)
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text,fmt).date().isoformat()
        except ValueError:
            pass
    return ""


def display_date(value):
    text=_text(value)
    try:
        return date.fromisoformat(text).strftime("%d.%m.%Y")
    except Exception:
        return text


def _date_from_note(value):
    match=re.search(r"(?<!\d)(\d{2}\.\d{2}\.\d{4})(?!\d)", _text(value))
    return _iso_date(match.group(1)) if match else ""


def _valid_rnokpp(value):
    text=re.sub(r"\s+", "", _text(value))
    return text if re.fullmatch(r"\d{10}", text) else ""


def _valid_id_card(value):
    text=re.sub(r"\s+", "", _text(value))
    return text if re.fullmatch(r"\d{9}", text) else ""


def _passport_parts(value):
    raw=_text(value)
    compact=re.sub(r"\s+", "", raw).upper()
    match=re.fullmatch(r"([A-ZА-ЯІЇЄҐ]{2})(\d{6})", compact)
    if match:
        return match.group(1), match.group(2)
    if compact and compact not in {"-", "—"}:
        return "", raw
    return "", ""


def _summary_passport(series, number):
    ser=re.sub(r"\s+", "", _text(series)).upper()
    num=re.sub(r"\s+", "", _text(number))
    if re.fullmatch(r"[A-ZА-ЯІЇЄҐ]{2}", ser) and re.fullmatch(r"\d{6}", num):
        return ser,num
    return "",""


def _header_map(values):
    return {_norm(value):idx for idx,value in enumerate(values) if _text(value)}


def detect_registry_form(headers):
    keys=set(_header_map(headers))
    if _DETAILED_HEADERS.issubset(keys):
        return SOURCE_DETAILED
    if _SUMMARY_HEADERS.issubset(keys):
        return SOURCE_SUMMARY
    return ""


def _value_by_header(row, header_index, header):
    idx=header_index.get(_norm(header))
    if idx is None or idx>=len(row):
        return ""
    return row[idx]


def _detailed_record(row, h):
    full_name=_text(_value_by_header(row,h,"Прізвище, ім’я, по батькові"))
    last,first,middle=_split_full_name(full_name)
    rnokpp=_valid_rnokpp(_value_by_header(row,h,"РНОКПП"))
    passport_series,passport_number=_passport_parts(
        _value_by_header(row,h,"Серія та номер паспорту")
    )
    id_card=_valid_id_card(_value_by_header(row,h,"Номер ID картки"))
    personal={
        "birth_date":_iso_date(_value_by_header(row,h,"Дата народження")),
        "rnokpp":rnokpp,
        "passport_series":passport_series,
        "passport_number":passport_number,
        "id_card_number":id_card,
        "registered_address":_text(_value_by_header(row,h,"Адреса зареєстрованого місця проживання")),
        "actual_address":_text(_value_by_header(row,h,"Адреса фактичного місця проживання")),
    }
    military={
        "registry_person_id":_text(_value_by_header(row,h,"Ідентифікатор військовозобов'язаного")),
        "account_type":_text(_value_by_header(row,h,"Вид обліку особи")),
        "account_status":_text(_value_by_header(row,h,"Статус обліку особи")),
        "reservist_status":_text(_value_by_header(row,h,"Належність до резервістів")),
        "booking_status":_text(_value_by_header(row,h,"Бронювання")),
        "booking_until":_iso_date(_value_by_header(row,h,"Дата кінця бронювання")),
        "deferment_reason":_text(_value_by_header(row,h,"Причина відстрочки")),
        "deferment_until":_iso_date(_value_by_header(row,h,"Дата закінчення відстрочки")),
        "military_rank":_text(_value_by_header(row,h,"Військове звання")),
        "military_specialty":_text(_value_by_header(row,h,"Військово-облікова спеціальність")),
        "account_authority":_text(_value_by_header(row,h,"Найменування ТЦК, органу СБУ, відповідного підрозділу розвідувального органу, в якому перебуває на військовому обліку")),
        "military_service_status":_text(_value_by_header(row,h,"Військова служба")),
    }
    documents=[]
    if passport_number:
        documents.append({"doc_type":"Паспорт громадянина України","series":passport_series,"number":passport_number})
    if id_card:
        documents.append({"doc_type":"ID-картка","series":"","number":id_card})
    return {
        "last_name":last,"first_name":first,"middle_name":middle,"full_name":full_name,
        "personal":personal,"military":military,"documents":documents,"warnings":[],
    }


def _summary_record(row, h):
    last=_text(_value_by_header(row,h,"Прізвище"))
    first=_text(_value_by_header(row,h,"Імʼя")) or _text(_value_by_header(row,h,"Ім’я"))
    middle=_text(_value_by_header(row,h,"По батькові"))
    status=_text(_value_by_header(row,h,"Статус"))
    note=_text(_value_by_header(row,h,"Примітка"))
    series,number=_summary_passport(
        _value_by_header(row,h,"Серія паспорту"),
        _value_by_header(row,h,"Номер паспорту"),
    )
    raw_id=_text(_value_by_header(row,h,"Номер ID-картки"))
    id_card=_valid_id_card(raw_id)
    warnings=[]
    if (_text(_value_by_header(row,h,"Серія паспорту")) or _text(_value_by_header(row,h,"Номер паспорту"))) and not number:
        warnings.append("Реквізити паспорта у зведеній формі не схожі на повний номер — у картку не перенесено")
    if raw_id and not id_card:
        warnings.append("Номер ID-картки у зведеній формі не схожий на 9-значний номер — у картку не перенесено")
    personal={
        "gender":_text(_value_by_header(row,h,"Стать")),
        "birth_date":_iso_date(_value_by_header(row,h,"Дата народження")),
        "rnokpp":_valid_rnokpp(_value_by_header(row,h,"РНОКПП")),
        "passport_series":series,
        "passport_number":number,
        "id_card_number":id_card,
    }
    military={"booking_status":status}
    until=_date_from_note(note)
    if status.lower().startswith("заброньовано") and until:
        military["booking_status"]="Заброньовано до певної дати"
        military["booking_until"]=until
    note_norm=_norm(note)
    if "відстроч" in note_norm:
        military["deferment_reason"]=note
    if "виключено з військового обліку" in note_norm:
        military["account_status"]="Виключено з військового обліку"
    elif "не перебуває на військовому обліку" in note_norm:
        military["account_status"]="Не на обліку"
    documents=[]
    if number:
        documents.append({"doc_type":"Паспорт громадянина України","series":series,"number":number})
    if id_card:
        documents.append({"doc_type":"ID-картка","series":"","number":id_card})
    full_name=" ".join(x for x in (last,first,middle) if x)
    return {
        "last_name":last,"first_name":first,"middle_name":middle,"full_name":full_name,
        "personal":personal,"military":military,"documents":documents,"warnings":warnings,
    }


def parse_registry_xlsx(path):
    """Parse either supported state-registry XLSX layout.

    The parser is driven by normalized headers, not by filenames or real people.
    """
    from openpyxl import load_workbook

    path=Path(path)
    wb=load_workbook(filename=str(path),read_only=True,data_only=True)
    try:
        for ws in wb.worksheets:
            rows=list(ws.iter_rows(values_only=True))
            header_row=None; source_kind=""; header_index={}
            for idx,row in enumerate(rows[:20]):
                kind=detect_registry_form(row)
                if kind:
                    header_row=idx; source_kind=kind; header_index=_header_map(row); break
            if header_row is None:
                continue
            parsed=[]
            for row_no,row in enumerate(rows[header_row+1:],start=header_row+2):
                if not any(_text(v) for v in row):
                    continue
                item=(_detailed_record(row,header_index) if source_kind==SOURCE_DETAILED
                      else _summary_record(row,header_index))
                item["source_row"]=row_no
                if not item["full_name"] and not item["personal"].get("rnokpp"):
                    continue
                parsed.append(item)
            return {
                "source_kind":source_kind,
                "sheet_name":ws.title,
                "header_row":header_row+1,
                "rows":parsed,
                "source_name":path.name,
            }
    finally:
        wb.close()
    raise ValueError("Формат XLSX не розпізнано як підтримуваний витяг з Реєстру.")


def file_sha256(path):
    digest=hashlib.sha256()
    with open(path,"rb") as fh:
        for chunk in iter(lambda:fh.read(1024*1024),b""):
            digest.update(chunk)
    return digest.hexdigest()


def ensure_schema_on_connection(con):
    employee_cols={row[1] for row in con.execute("PRAGMA table_info(employees)").fetchall()}
    for name,ddl in (
        ("gender","TEXT DEFAULT ''"),
        ("birth_date","TEXT DEFAULT ''"),
        ("rnokpp","TEXT DEFAULT ''"),
        ("passport_series","TEXT DEFAULT ''"),
        ("passport_number","TEXT DEFAULT ''"),
        ("id_card_number","TEXT DEFAULT ''"),
        ("registered_address","TEXT DEFAULT ''"),
        ("actual_address","TEXT DEFAULT ''"),
    ):
        if name not in employee_cols:
            con.execute("ALTER TABLE employees ADD COLUMN %s %s" % (name,ddl))
    con.execute("CREATE INDEX IF NOT EXISTS idx_employees_rnokpp ON employees(rnokpp)")
    con.executescript("""
        CREATE TABLE IF NOT EXISTS employee_military_profile (
            employee_id INTEGER PRIMARY KEY REFERENCES employees(id) ON DELETE CASCADE,
            registry_person_id TEXT DEFAULT '',
            account_type TEXT DEFAULT '',
            account_status TEXT DEFAULT '',
            reservist_status TEXT DEFAULT '',
            booking_status TEXT DEFAULT '',
            booking_until TEXT DEFAULT '',
            deferment_reason TEXT DEFAULT '',
            deferment_until TEXT DEFAULT '',
            military_rank TEXT DEFAULT '',
            military_specialty TEXT DEFAULT '',
            account_authority TEXT DEFAULT '',
            military_service_status TEXT DEFAULT '',
            reserve_category TEXT DEFAULT '',
            military_document TEXT DEFAULT '',
            education_specialty TEXT DEFAULT '',
            foreign_passport TEXT DEFAULT '',
            fitness TEXT DEFAULT '',
            family_info TEXT DEFAULT '',
            appointment_act TEXT DEFAULT '',
            notification_requisites TEXT DEFAULT '',
            last_source_kind TEXT DEFAULT '',
            last_source_name TEXT DEFAULT '',
            last_verified_at TEXT DEFAULT '',
            updated_at TEXT DEFAULT ''
        );
        CREATE TABLE IF NOT EXISTS employee_registry_imports (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            source_kind TEXT NOT NULL,
            source_name TEXT NOT NULL,
            file_sha256 TEXT NOT NULL,
            imported_at TEXT NOT NULL,
            total_rows INTEGER NOT NULL DEFAULT 0,
            matched_rows INTEGER NOT NULL DEFAULT 0,
            updated_rows INTEGER NOT NULL DEFAULT 0,
            skipped_rows INTEGER NOT NULL DEFAULT 0,
            warning_count INTEGER NOT NULL DEFAULT 0
        );
        CREATE INDEX IF NOT EXISTS idx_employee_registry_imports_date
            ON employee_registry_imports(imported_at DESC);
        CREATE TABLE IF NOT EXISTS employee_military_history (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            employee_id INTEGER NOT NULL REFERENCES employees(id) ON DELETE CASCADE,
            import_id INTEGER REFERENCES employee_registry_imports(id) ON DELETE SET NULL,
            source_kind TEXT NOT NULL,
            source_name TEXT NOT NULL,
            file_sha256 TEXT NOT NULL,
            source_row INTEGER,
            row_fingerprint TEXT NOT NULL,
            snapshot_json TEXT NOT NULL,
            imported_at TEXT NOT NULL,
            UNIQUE(employee_id,row_fingerprint)
        );
        CREATE INDEX IF NOT EXISTS idx_employee_military_history_employee
            ON employee_military_history(employee_id,imported_at DESC);
        CREATE TABLE IF NOT EXISTS employee_documents (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            employee_id INTEGER NOT NULL REFERENCES employees(id) ON DELETE CASCADE,
            doc_type TEXT NOT NULL,
            series TEXT DEFAULT '',
            number TEXT DEFAULT '',
            issue_date TEXT DEFAULT '',
            expiry_date TEXT DEFAULT '',
            issuer TEXT DEFAULT '',
            source_kind TEXT DEFAULT '',
            source_name TEXT DEFAULT '',
            file_path TEXT DEFAULT '',
            notes TEXT DEFAULT '',
            active INTEGER NOT NULL DEFAULT 1,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        );
        CREATE INDEX IF NOT EXISTS idx_employee_documents_employee
            ON employee_documents(employee_id,active,doc_type);
    """)
    import_cols={row[1] for row in con.execute("PRAGMA table_info(employee_registry_imports)").fetchall()}
    if "mode" not in import_cols:
        con.execute("ALTER TABLE employee_registry_imports ADD COLUMN mode TEXT NOT NULL DEFAULT 'update'")


def ensure_schema(core):
    con=core.db()
    try:
        ensure_schema_on_connection(con)
        con.commit()
    finally:
        con.close()


def _military_row(con, employee_id):
    return con.execute(
        "SELECT * FROM employee_military_profile WHERE employee_id=?",(employee_id,)
    ).fetchone()


def _employee_dict(row):
    return {key:row[key] for key in row.keys()}


def _match_employee(con, item):
    rnokpp=item["personal"].get("rnokpp","")
    employees=con.execute("SELECT * FROM employees ORDER BY active DESC,id").fetchall()
    if rnokpp:
        by_tax=[row for row in employees if _text(row["rnokpp"])==rnokpp]
        if len(by_tax)==1:
            return by_tax[0],"rnokpp",""
        if len(by_tax)>1:
            return None,"ambiguous","У реєстрі працівників є кілька однакових РНОКПП"
    key=_name_key(item["last_name"],item["first_name"],item["middle_name"])
    by_name=[row for row in employees if _name_key(row["last_name"],row["first_name"],row["middle_name"])==key]
    incoming_birth=item["personal"].get("birth_date","")
    if len(by_name)>1 and incoming_birth:
        by_birth=[row for row in by_name if _text(row["birth_date"])==incoming_birth]
        if len(by_birth)==1:
            by_name=by_birth
    if len(by_name)==1:
        row=by_name[0]
        existing_tax=_text(row["rnokpp"])
        if rnokpp and existing_tax and existing_tax!=rnokpp:
            return None,"conflict","ПІБ збігається, але РНОКПП у картці інший"
        existing_birth=_text(row["birth_date"])
        if incoming_birth and existing_birth and existing_birth!=incoming_birth:
            return None,"conflict","ПІБ збігається, але дата народження у картці інша"
        return row,"name",""
    if len(by_name)>1:
        return None,"ambiguous","ПІБ відповідає кільком карткам працівників"
    return None,"unmatched","Працівника не знайдено — автоматично нову картку не створено"


def _change(scope, field, old, new):
    labels=EMPLOYEE_FIELD_LABELS if scope=="employee" else MILITARY_FIELD_LABELS
    return {"scope":scope,"field":field,"label":labels.get(field,field),"old":_text(old),"new":_text(new)}


def _specificity_guard(source_kind, field, old, new):
    if source_kind==SOURCE_SUMMARY and field=="booking_status":
        if _text(old)=="Заброньовано до певної дати" and _text(new)=="Заброньовано":
            return False
    return True


def plan_import_rows(con, parsed):
    ensure_schema_on_connection(con)
    plan=[]
    for item in parsed["rows"]:
        employee,match_quality,note=_match_employee(con,item)
        if employee is None:
            plan.append({"item":item,"employee_id":None,"employee_name":item["full_name"],
                         "match_quality":match_quality,"status":match_quality,"changes":[],
                         "notes":[note]+list(item.get("warnings") or [])})
            continue
        employee_data=_employee_dict(employee)
        military=_military_row(con,employee["id"])
        military_data=_employee_dict(military) if military is not None else {}
        changes=[]
        for field,new in item["personal"].items():
            new=_text(new)
            if not new:
                continue
            old=_text(employee_data.get(field,""))
            if old!=new:
                changes.append(_change("employee",field,old,new))
        for field,new in item["military"].items():
            new=_text(new)
            if not new:
                continue
            old=_text(military_data.get(field,""))
            if old!=new and _specificity_guard(parsed["source_kind"],field,old,new):
                changes.append(_change("military",field,old,new))
        status="update" if changes else "no_changes"
        plan.append({
            "item":item,"employee_id":employee["id"],
            "employee_name":" ".join(x for x in (employee["last_name"],employee["first_name"],employee["middle_name"]) if x),
            "match_quality":match_quality,"status":status,"changes":changes,
            "notes":list(item.get("warnings") or []),
        })
    return plan


def preview_registry_import(core, path):
    parsed=parse_registry_xlsx(path)
    sha=file_sha256(path)
    con=core.db()
    try:
        plan=plan_import_rows(con,parsed)
        matched_ids={int(row["employee_id"]) for row in plan if row.get("employee_id") is not None}
        local_only=[]
        for employee in con.execute("SELECT * FROM employees WHERE COALESCE(active,1)=1 ORDER BY last_name,first_name,middle_name,id").fetchall():
            if int(employee["id"]) in matched_ids:
                continue
            local_only.append({
                "employee_id":int(employee["id"]),
                "employee_name":" ".join(x for x in (employee["last_name"],employee["first_name"],employee["middle_name"]) if x),
                "status":"local_only",
                "note":"Є у Taxo, але відсутній у цьому витягу. Дані не видаляються і не обнуляються.",
            })
    finally:
        con.close()
    return {"path":str(path),"file_sha256":sha,"parsed":parsed,"plan":plan,"local_only":local_only}


def _changes_for_mode(changes, mode):
    if mode not in IMPORT_MODES:
        raise ValueError("Невідомий режим реєстрової звірки: %s" % mode)
    if mode==IMPORT_COMPARE:
        return []
    if mode==IMPORT_FILL_EMPTY:
        return [change for change in changes if not _text(change.get("old")) and _text(change.get("new"))]
    return [change for change in changes if _text(change.get("new"))]


def registry_quarter_status(con, source_kind=None, today=None):
    """Return whether at least one successful reconciliation exists in current calendar quarter."""
    ensure_schema_on_connection(con)
    today=today or date.today()
    quarter=(today.month-1)//3+1
    first_month=(quarter-1)*3+1
    quarter_start=date(today.year,first_month,1)
    params=[]
    sql="SELECT imported_at,source_kind,source_name,mode FROM employee_registry_imports"
    if source_kind:
        sql += " WHERE source_kind=?"
        params.append(source_kind)
    sql += " ORDER BY imported_at DESC,id DESC LIMIT 1"
    row=con.execute(sql,params).fetchone()
    last_date=None
    if row is not None:
        try: last_date=datetime.fromisoformat(str(row["imported_at"])).date()
        except Exception: last_date=None
    current=bool(last_date and last_date>=quarter_start and last_date<=today)
    return {
        "current":current,
        "quarter":"Q%d %d" % (quarter,today.year),
        "label":"Актуально" if current else "Потрібне звіряння",
        "last_date":last_date.isoformat() if last_date else "",
        "source_kind":row["source_kind"] if row is not None else "",
        "source_name":row["source_name"] if row is not None else "",
        "mode":row["mode"] if row is not None and "mode" in row.keys() else "",
    }


def _canonical_snapshot(item):
    return {
        "personal":{k:_text(v) for k,v in sorted(item.get("personal",{}).items()) if _text(v)},
        "military":{k:_text(v) for k,v in sorted(item.get("military",{}).items()) if _text(v)},
        "documents":[{k:_text(v) for k,v in sorted(doc.items()) if _text(v)} for doc in item.get("documents",[])],
    }


def _row_fingerprint(item):
    payload=json.dumps(_canonical_snapshot(item),ensure_ascii=False,sort_keys=True,separators=(",",":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _ensure_profile(con, employee_id):
    con.execute("INSERT OR IGNORE INTO employee_military_profile(employee_id) VALUES(?)",(employee_id,))


def _upsert_document_hint(con, employee_id, doc, source_kind, source_name, now):
    doc_type=_text(doc.get("doc_type")); series=_text(doc.get("series")); number=_text(doc.get("number"))
    if not doc_type or not number:
        return
    existing=con.execute(
        """SELECT id FROM employee_documents
           WHERE employee_id=? AND doc_type=? AND COALESCE(series,'')=? AND COALESCE(number,'')=?
           ORDER BY active DESC,id LIMIT 1""",
        (employee_id,doc_type,series,number),
    ).fetchone()
    if existing:
        con.execute(
            """UPDATE employee_documents SET active=1,source_kind=?,source_name=?,updated_at=?
               WHERE id=?""",(source_kind,source_name,now,existing["id"])
        )
    else:
        con.execute(
            """INSERT INTO employee_documents(
                   employee_id,doc_type,series,number,source_kind,source_name,active,created_at,updated_at
               ) VALUES(?,?,?,?,?,?,1,?,?)""",
            (employee_id,doc_type,series,number,source_kind,source_name,now,now),
        )


def apply_registry_import(core, preview, mode=IMPORT_UPDATE):
    if mode not in IMPORT_MODES:
        raise ValueError("Невідомий режим реєстрової звірки: %s" % mode)
    path=preview["path"]
    current_sha=file_sha256(path)
    if current_sha!=preview["file_sha256"]:
        raise ValueError("Файл змінився після попереднього перегляду. Виконайте імпорт ще раз.")
    parsed=preview["parsed"]
    plan=preview["plan"]
    now=datetime.now().isoformat(timespec="seconds")
    con=core.db()
    try:
        ensure_schema_on_connection(con)
        cur=con.execute(
            """INSERT INTO employee_registry_imports(
                   source_kind,source_name,file_sha256,imported_at,total_rows,mode
               ) VALUES(?,?,?,?,?,?)""",
            (parsed["source_kind"],parsed["source_name"],current_sha,now,len(plan),mode),
        )
        import_id=cur.lastrowid
        matched=updated=skipped=warnings=0
        potential_changes=0
        for row_plan in plan:
            warnings += len(row_plan.get("notes") or [])
            employee_id=row_plan.get("employee_id")
            if employee_id is None:
                skipped += 1
                continue
            matched += 1
            item=row_plan["item"]
            potential_changes += len(row_plan.get("changes") or [])
            selected_changes=_changes_for_mode(row_plan.get("changes") or [],mode)
            employee_changes=[c for c in selected_changes if c["scope"]=="employee"]
            if employee_changes:
                sql="UPDATE employees SET "+", ".join(c["field"]+"=?" for c in employee_changes)+" WHERE id=?"
                con.execute(sql,[c["new"] for c in employee_changes]+[employee_id])
            military_changes=[c for c in selected_changes if c["scope"]=="military"]
            _ensure_profile(con,employee_id)
            if military_changes:
                sql="UPDATE employee_military_profile SET "+", ".join(c["field"]+"=?" for c in military_changes)+", last_source_kind=?, last_source_name=?, last_verified_at=?, updated_at=? WHERE employee_id=?"
                con.execute(sql,[c["new"] for c in military_changes]+[
                    parsed["source_kind"],parsed["source_name"],now,now,employee_id
                ])
            else:
                con.execute(
                    """UPDATE employee_military_profile SET last_source_kind=?,last_source_name=?,last_verified_at=?,updated_at=?
                       WHERE employee_id=?""",
                    (parsed["source_kind"],parsed["source_name"],now,now,employee_id),
                )
            if selected_changes:
                updated += 1
            snapshot=_canonical_snapshot(item)
            fingerprint=_row_fingerprint(item)
            con.execute(
                """INSERT OR IGNORE INTO employee_military_history(
                       employee_id,import_id,source_kind,source_name,file_sha256,source_row,
                       row_fingerprint,snapshot_json,imported_at
                   ) VALUES(?,?,?,?,?,?,?,?,?)""",
                (employee_id,import_id,parsed["source_kind"],parsed["source_name"],current_sha,
                 item.get("source_row"),fingerprint,
                 json.dumps(snapshot,ensure_ascii=False,sort_keys=True),now),
            )
            # Compare-only must never mutate the local document register.
            if mode!=IMPORT_COMPARE:
                for doc in item.get("documents",[]):
                    _upsert_document_hint(con,employee_id,doc,parsed["source_kind"],parsed["source_name"],now)
        con.execute(
            """UPDATE employee_registry_imports
               SET matched_rows=?,updated_rows=?,skipped_rows=?,warning_count=? WHERE id=?""",
            (matched,updated,skipped,warnings,import_id),
        )
        con.commit()
        return {"import_id":import_id,"total":len(plan),"matched":matched,"updated":updated,
                "skipped":skipped,"warnings":warnings,"potential_changes":potential_changes,
                "local_only":len(preview.get("local_only") or []),"mode":mode,
                "source_kind":parsed["source_kind"],"source_name":parsed["source_name"]}
    except Exception:
        con.rollback()
        raise
    finally:
        con.close()


def employee_profile(con, employee_id):
    ensure_schema_on_connection(con)
    employee=con.execute("SELECT * FROM employees WHERE id=?",(employee_id,)).fetchone()
    military=_military_row(con,employee_id)
    return employee,military


def missing_appendix5_fields(employee, military):
    missing=[]
    if employee is None:
        return missing
    employee_keys=set(employee.keys())
    for field in ("birth_date","rnokpp","passport_number","registered_address","actual_address"):
        if field in employee_keys and not _text(employee[field]):
            missing.append(EMPLOYEE_FIELD_LABELS.get(field,field))
    data=_employee_dict(military) if military is not None else {}
    for field in APPENDIX5_PROFILE_FIELDS:
        if not _text(data.get(field,"")):
            missing.append(MILITARY_FIELD_LABELS.get(field,field))
    return missing


def list_employee_documents(con, employee_id, include_archived=False):
    ensure_schema_on_connection(con)
    sql="SELECT * FROM employee_documents WHERE employee_id=?"
    params=[employee_id]
    if not include_archived:
        sql += " AND active=1"
    sql += " ORDER BY active DESC,doc_type,expiry_date DESC,id DESC"
    return con.execute(sql,params).fetchall()


def save_employee_document(con, employee_id, values, document_id=None):
    ensure_schema_on_connection(con)
    now=datetime.now().isoformat(timespec="seconds")
    doc_type=_text(values.get("doc_type"))
    if not doc_type:
        raise ValueError("Вкажіть тип документа.")
    payload={
        "doc_type":doc_type,
        "series":_text(values.get("series")),
        "number":_text(values.get("number")),
        "issue_date":_iso_date(values.get("issue_date")) or _text(values.get("issue_date")),
        "expiry_date":_iso_date(values.get("expiry_date")) or _text(values.get("expiry_date")),
        "issuer":_text(values.get("issuer")),
        "source_kind":_text(values.get("source_kind")) or "manual",
        "source_name":_text(values.get("source_name")),
        "file_path":_text(values.get("file_path")),
        "notes":_text(values.get("notes")),
    }
    if document_id:
        con.execute(
            """UPDATE employee_documents SET doc_type=?,series=?,number=?,issue_date=?,expiry_date=?,
                   issuer=?,source_kind=?,source_name=?,file_path=?,notes=?,updated_at=?
               WHERE id=? AND employee_id=?""",
            (payload["doc_type"],payload["series"],payload["number"],payload["issue_date"],
             payload["expiry_date"],payload["issuer"],payload["source_kind"],payload["source_name"],
             payload["file_path"],payload["notes"],now,document_id,employee_id),
        )
        return document_id
    cur=con.execute(
        """INSERT INTO employee_documents(
               employee_id,doc_type,series,number,issue_date,expiry_date,issuer,source_kind,
               source_name,file_path,notes,active,created_at,updated_at
           ) VALUES(?,?,?,?,?,?,?,?,?,?,?,1,?,?)""",
        (employee_id,payload["doc_type"],payload["series"],payload["number"],payload["issue_date"],
         payload["expiry_date"],payload["issuer"],payload["source_kind"],payload["source_name"],
         payload["file_path"],payload["notes"],now,now),
    )
    return cur.lastrowid


def archive_employee_document(con, employee_id, document_id):
    con.execute(
        "UPDATE employee_documents SET active=0,updated_at=? WHERE id=? AND employee_id=?",
        (datetime.now().isoformat(timespec="seconds"),document_id,employee_id),
    )


def list_import_history(con, employee_id=None, limit=200):
    ensure_schema_on_connection(con)
    if employee_id is None:
        return con.execute(
            "SELECT * FROM employee_registry_imports ORDER BY imported_at DESC,id DESC LIMIT ?",(int(limit),)
        ).fetchall()
    return con.execute(
        """SELECT h.*,i.total_rows,i.matched_rows,i.updated_rows,i.skipped_rows
             FROM employee_military_history h
        LEFT JOIN employee_registry_imports i ON i.id=h.import_id
            WHERE h.employee_id=? ORDER BY h.imported_at DESC,h.id DESC LIMIT ?""",
        (employee_id,int(limit)),
    ).fetchall()
