import tempfile
import unittest
from datetime import date
from pathlib import Path


_temporary_home = tempfile.TemporaryDirectory()
_original_path_home = Path.home
Path.home = classmethod(lambda cls: Path(_temporary_home.name))
try:
    import fitz
    from openpyxl import load_workbook
    import main
finally:
    Path.home = _original_path_home


class V870R8Tests(unittest.TestCase):
    def setUp(self):
        main.init_db()

    def _employee(self, con, suffix="r8", employment="2031-09-05", dismissal="2031-09-20"):
        now="2031-09-01T10:00:00"
        return con.execute(
            """INSERT INTO employees(personnel_no,last_name,first_name,position,employment_date,dismissal_date,active,created_at)
               VALUES(?,?,?,?,?,?,1,?)""",
            (f"R8-{suffix}",f"Працівник-{suffix}","Тестовий","Механік",employment,dismissal,now),
        ).lastrowid

    def test_employee_collection_respects_employment_and_plan_fact(self):
        con=main.db()
        try:
            employee_id=self._employee(con,"collect")
            con.execute("""INSERT INTO employee_shifts(employee_id,role,work_date,shift_no,start_time,end_day_offset,end_time,planned_hours,status)
                           VALUES(?,?,?,?,?,?,?,?,?)""",
                        (employee_id,"Механік","2031-09-15",1,"08:00",0,"16:00",8.0,"planned"))
            now="2031-09-15T18:00:00"
            con.execute("""INSERT INTO employee_time_entries(employee_id,work_date,day_type,planned_hours,actual_hours,notes,created_at,updated_at)
                           VALUES(?,?,?,?,?,?,?,?)""",
                        (employee_id,"2031-09-15","Робота",None,7.5,"факт r8",now,now))
            con.commit()
        finally:
            con.close()
        data=main.collect_employee_timesheet(employee_id,2031,9)
        by_date={row["date"]:row for row in data["rows"]}
        self.assertFalse(by_date[date(2031,9,1)]["employed"])
        self.assertEqual(by_date[date(2031,9,15)]["planned_minutes"],480)
        self.assertEqual(by_date[date(2031,9,15)]["actual_minutes"],450)
        self.assertEqual(data["difference_minutes"],-30)

    def test_personnel_balance_marks_missing_fact(self):
        con=main.db()
        try:
            employee_id=self._employee(con,"missing",employment="2031-09-01",dismissal="")
            now="2031-09-10T10:00:00"
            con.execute("""INSERT INTO employee_time_entries(employee_id,work_date,day_type,planned_hours,actual_hours,notes,created_at,updated_at)
                           VALUES(?,?,?,?,?,?,?,?)""",
                        (employee_id,"2031-09-10","Робота",8.0,None,"",now,now))
            con.commit()
        finally:
            con.close()
        data=main.collect_personnel_monthly_balance(2031,9,active_only=True)
        row=next(item for item in data["employees"] if item["employee_id"]==employee_id)
        self.assertEqual(row["cells"][9],"—")
        self.assertEqual(row["missing_days"],1)

    def test_employee_xlsx_and_pdf_exports(self):
        con=main.db()
        try:
            employee_id=self._employee(con,"export",employment="2031-09-01",dismissal="")
            now="2031-09-12T10:00:00"
            con.execute("""INSERT INTO employee_time_entries(employee_id,work_date,day_type,planned_hours,actual_hours,notes,created_at,updated_at)
                           VALUES(?,?,?,?,?,?,?,?)""",
                        (employee_id,"2031-09-12","Робота",8.0,8.0,"зміна",now,now))
            con.commit()
        finally:
            con.close()
        with tempfile.TemporaryDirectory() as folder:
            xlsx=Path(folder)/"employee.xlsx"; pdf=Path(folder)/"employee.pdf"
            main.export_employee_timesheet_xlsx(employee_id,2031,9,xlsx)
            main.export_employee_timesheet_pdf(employee_id,2031,9,pdf)
            wb=load_workbook(xlsx,data_only=False)
            self.assertEqual(wb.sheetnames,["Щоденний табель"])
            self.assertEqual(wb["Щоденний табель"]["A4"].value,"Дата")
            doc=fitz.open(pdf); text="\n".join(page.get_text() for page in doc); doc.close()
            self.assertIn("ТАБЕЛЬ РОБОЧОГО ЧАСУ",text)
            self.assertIn("Працівник-export",text)

    def test_personnel_balance_xlsx_and_pdf_exports(self):
        con=main.db()
        try:
            employee_id=self._employee(con,"balance",employment="2031-09-01",dismissal="")
            now="2031-09-08T10:00:00"
            con.execute("""INSERT INTO employee_time_entries(employee_id,work_date,day_type,planned_hours,actual_hours,notes,created_at,updated_at)
                           VALUES(?,?,?,?,?,?,?,?)""",
                        (employee_id,"2031-09-08","Робота",8.0,7.5,"",now,now))
            con.commit()
        finally:
            con.close()
        with tempfile.TemporaryDirectory() as folder:
            xlsx=Path(folder)/"personnel.xlsx"; pdf=Path(folder)/"personnel.pdf"
            main.export_personnel_monthly_balance_xlsx(2031,9,xlsx)
            main.export_personnel_monthly_balance_pdf(2031,9,pdf)
            wb=load_workbook(xlsx,data_only=False)
            self.assertGreaterEqual(len(wb.sheetnames),2)
            self.assertEqual(wb[wb.sheetnames[0]]["A4"].value,"Таб. №")
            doc=fitz.open(pdf); text="\n".join(page.get_text() for page in doc); doc.close()
            self.assertIn("ТАБЕЛЬ УСЬОГО ПЕРСОНАЛУ",text)

    def test_r8_release_workflow_builds_every_platform_and_checks_databases(self):
        root=Path(__file__).resolve().parents[1]
        workflow=(root/".github/workflows/publish-v8.70.yml").read_text("utf-8")
        self.assertIn("Taxo_v8_70_TEST_r8_Setup_Windows_x64.exe",workflow)
        self.assertIn("Taxo_v8_70_TEST_r8_Windows_x64_Portable.zip",workflow)
        self.assertIn("Taxo_v8_70_TEST_r8_macOS_arm64_Portable.zip",workflow)
        self.assertIn("Taxo_v8_70_TEST_r8_macOS_x86_64_Portable.zip",workflow)
        self.assertIn("SHA256SUMS_v8_70_TEST_r8.txt",workflow)
        self.assertIn("Database unexpectedly bundled",workflow)
        self.assertIn("needs:",workflow)


if __name__ == "__main__":
    unittest.main()
