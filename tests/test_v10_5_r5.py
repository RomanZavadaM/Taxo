# -*- coding: utf-8 -*-
import tempfile
import unittest
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import Workbook, load_workbook

import vehicle_registry
import v1055_features

ROOT = Path(__file__).resolve().parents[1]

HEADERS = [
    "Вид", "№ ТЗ", "Перевізник", "Статус", "VIN код автомобіля", "Марка", "Модель",
]
ROW = [
    "Автобус", "BC1234AA", "12345678 Тестовий перевізник", "На обліку",
    "Y6A00000000000001", "Еталон", "А-079",
]


def make_shlyakh_xlsx(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Транспортні засоби"
    ws.append(HEADERS)
    ws.append(ROW)
    wb.save(path)
    wb.close()


def inject_sheet_tabid(path, value="123"):
    source = Path(path)
    rewritten = source.with_suffix(".rewrite.xlsx")
    with zipfile.ZipFile(str(source), "r") as src, zipfile.ZipFile(str(rewritten), "w") as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "xl/workbook.xml":
                root = ET.fromstring(data)
                sheets = [elem for elem in root.iter() if elem.tag.rsplit("}", 1)[-1] == "sheet"]
                if not sheets:
                    raise AssertionError("Synthetic XLSX has no <sheet> element")
                sheets[0].set("tabId", value)
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            dst.writestr(info, data)
    rewritten.replace(source)


class ShlyakhTabIdCompatibilityR5Tests(unittest.TestCase):
    def test_fixture_reproduces_openpyxl_childsheet_tabid_error(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "shlyakh-tabid.xlsx"
            make_shlyakh_xlsx(path)
            inject_sheet_tabid(path)
            with self.assertRaises(TypeError) as cm:
                load_workbook(filename=str(path), read_only=True, data_only=True)
            self.assertIn("ChildSheet", str(cm.exception))
            self.assertIn("tabId", str(cm.exception))

    def test_taxo_reads_tabid_export_without_modifying_source(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "shlyakh-tabid.xlsx"
            make_shlyakh_xlsx(path)
            inject_sheet_tabid(path)
            before = path.read_bytes()

            old_parser = vehicle_registry.parse_registry_xlsx
            try:
                vehicle_registry.parse_registry_xlsx = v1055_features.parse_registry_xlsx_compat
                parsed = vehicle_registry.parse_registry_file(path)
            finally:
                vehicle_registry.parse_registry_xlsx = old_parser

            self.assertEqual(before, path.read_bytes(), "Taxo must never rewrite the user's registry XLSX")
            self.assertEqual(parsed["source_kind"], vehicle_registry.SOURCE_KIND)
            self.assertEqual(parsed["sheet_name"], "Транспортні засоби")
            self.assertEqual(len(parsed["rows"]), 1)
            self.assertEqual(parsed["rows"][0]["plate"], "BC1234AA")
            self.assertEqual(parsed["rows"][0]["vin"], "Y6A00000000000001")
            self.assertEqual(parsed["rows"][0]["make"], "Еталон")
            self.assertEqual(parsed["rows"][0]["model"], "А-079")

    def test_normal_xlsx_still_uses_same_parser_contract(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "shlyakh-normal.xlsx"
            make_shlyakh_xlsx(path)
            parsed = v1055_features.parse_registry_xlsx_compat(path)
            self.assertEqual(len(parsed["rows"]), 1)
            self.assertEqual(parsed["rows"][0]["plate"], "BC1234AA")

    def test_unrelated_typeerror_is_not_classified_as_tabid_compatibility_case(self):
        self.assertFalse(v1055_features._is_childsheet_tabid_error(TypeError("other failure")))
        self.assertTrue(v1055_features._is_childsheet_tabid_error(
            TypeError("ChildSheet.__init__() got an unexpected keyword argument 'tabId'")
        ))


class R5IdentityTests(unittest.TestCase):
    def test_candidate_identity_is_r5(self):
        self.assertEqual(v1055_features.APP_VERSION, "10.5-r5")
        self.assertIn("Version: 10.5-r5", (ROOT / "VERSION.txt").read_text(encoding="utf-8"))

    def test_r5_layer_is_outermost_entrypoint(self):
        text = (ROOT / "taxo_app.py").read_text(encoding="utf-8")
        self.assertIn("from v1055_features import install as install_v1055", text)
        self.assertIn("App = install_v1055(", text)


if __name__ == "__main__":
    unittest.main()
