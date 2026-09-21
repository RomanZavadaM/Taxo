# -*- coding: utf-8 -*-
import calendar
import sqlite3
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace

from personnel_v91 import (
    ABSENCE_TYPES,
    NONWORK_OVERRIDE_TYPES,
    P5_CODES,
    P5_LEGEND,
    collect_p5_data,
    export_p5_pdf,
    export_p5_xlsx,
    install,
    p5_code,
    driver_plan_conflict,
    _legacy_absence_cell,
)


class FakeCore:
    def __init__(self, db_path):
        self.DB_PATH = Path(db_path)
        self.App = object
        self._TAXO_PERSONNEL_R3_INSTALLED = False

    def db(self):
        con = sqlite3.connect(self.DB_PATH)
        con.row_factory = sqlite3.Row
        return con

    @staticmethod
    def hours_value_to_minutes(value):
        if value is None:
            return 0
        return int(round(float(value) * 60))

    @staticmethod
    def duration_minutes(start, end):
        sh, sm = map(int, start.split(":"))
        eh, em = map(int, end.split(":"))
        a = sh * 60 + sm
        b = eh * 60 + em
        if b <= a:
            b += 1440
        return b - a

    @staticmethod
    def minutes_hhmm(value):
        value = int(value or 0)
        return f"{value // 60}:{value % 60:02d}"

    @staticmethod
    def month_dates(year, month):
        last = calendar.monthrange(int(year), int(month))[1]
        return [date(int(year), int(month), day) for day in range(1, last + 1)]

    @staticmethod
    def employee_employed_on(employee, target_date):
        started = date.fromisoformat(employee["employment_date"]) if employee["employment_date"] else None
        finished = date.fromisoformat(employee["dismissal_date"]) if employee["dismissal_date"] else None
        return (started is None or target_date >= started) and (finished is None or target_date <= finished)

    @staticmethod
    def employee_name(employee):
        return " ".join(x for x in (employee["last_name"], employee["first_name"], employee["middle_name"]) if x)

    @staticmethod
    def month_name_ua(month):
        return ("Січень","Лютий","Березень","Квітень","Травень","Червень","Липень","Серпень","Вересень","Жовтень","Листопад","Грудень")[int(month)-1]

    @staticmethod
    def report_font_candidates():
        return []

    def collect_monthly_work_balance(self, year, month, active_only=True):
        days = self.month_dates(year, month)
        return {
            "year": int(year),
            "month": int(month),
            "days": days,
            "drivers": [{
                "driver_id": 10,
                "personnel_no": "D10",
                "name": "Тестовий Водій",
                "cells": ["8:00"] * len(days),
                "work_days": len(days),
                "work_min": len(days) * 480,
                "over_min": 0,
            }],
            "company": None,
        }

    def employee_day_time(self, con, employee_id, target_date):
        entry = con.execute(
            "SELECT * FROM employee_time_entries WHERE employee_id=? AND work_date=?",
            (employee_id, target_date.isoformat()),
        ).fetchone()
        if entry:
            planned = self.hours_value_to_minutes(entry["planned_hours"]) if entry["planned_hours"] is not None else 0
            actual = self.hours_value_to_minutes(entry["actual_hours"]) if entry["actual_hours"] is not None else None
            return {
                "day_type": entry["day_type"],
                "planned_minutes": planned,
                "actual_minutes": actual,
                "source": "manual",
                "notes": entry["notes"] or "",
                "manual": True,
            }
        return {
            "day_type": "Вихідний",
            "planned_minutes": 0,
            "actual_minutes": None,
            "source": "—",
            "notes": "",
            "manual": False,
        }


def make_db(path):
    con = sqlite3.connect(path)
    con.executescript(
        """
        CREATE TABLE company(id INTEGER PRIMARY KEY, name TEXT);
        INSERT INTO company(id,name) VALUES(1,'Тестове АТП');
        CREATE TABLE employees(
            id INTEGER PRIMARY KEY,
            personnel_no TEXT,
            last_name TEXT,
            first_name TEXT,
            middle_name TEXT,
            position TEXT,
            employment_date TEXT,
            dismissal_date TEXT,
            active INTEGER,
            driver_id INTEGER
        );
        CREATE TABLE employee_roles(employee_id INTEGER, role TEXT);
        CREATE TABLE employee_time_entries(
            id INTEGER PRIMARY KEY,
            employee_id INTEGER,
            work_date TEXT,
            day_type TEXT,
            planned_hours REAL,
            actual_hours REAL,
            notes TEXT
        );
        CREATE TABLE worklog(
            id INTEGER PRIMARY KEY,
            driver_id INTEGER,
            work_date TEXT,
            overtime_hours REAL,
            route_id INTEGER,
            work_hours REAL DEFAULT 0,
            work_start_time TEXT DEFAULT '',
            work_end_time TEXT DEFAULT ''
        );
        CREATE TABLE routes(
            id INTEGER PRIMARY KEY,
            name TEXT
        );
        CREATE TABLE route_segments(
            id INTEGER PRIMARY KEY,
            route_id INTEGER,
            segment_no INTEGER,
            start_time TEXT DEFAULT '',
            end_time TEXT DEFAULT '',
            work_start_time TEXT DEFAULT '',
            work_end_time TEXT DEFAULT '',
            work_hours REAL DEFAULT 0,
            driving_hours REAL DEFAULT 0
        );
        CREATE TABLE work_segments(
            id INTEGER PRIMARY KEY,
            worklog_id INTEGER,
            segment_no INTEGER,
            start_time TEXT DEFAULT '',
            end_time TEXT DEFAULT '',
            work_start_time TEXT DEFAULT '',
            work_end_time TEXT DEFAULT '',
            work_hours REAL DEFAULT 0,
            driving_hours REAL DEFAULT 0
        );
        INSERT INTO employees VALUES(1,'0001','Тестовий','Працівник','','механік','2026-01-01','',1,NULL);
        INSERT INTO employee_roles VALUES(1,'Механік');
        INSERT INTO employee_time_entries(employee_id,work_date,day_type,planned_hours,actual_hours,notes)
            VALUES(1,'2026-08-03','Робота',8,8,'');
        INSERT INTO employee_time_entries(employee_id,work_date,day_type,planned_hours,actual_hours,notes)
            VALUES(1,'2026-08-04','Основна щорічна відпустка',0,0,'наказ');
        INSERT INTO employee_time_entries(employee_id,work_date,day_type,planned_hours,actual_hours,notes)
            VALUES(1,'2026-08-05','Оплачувана тимчасова непрацездатність',0,0,'листок');
        """
    )
    con.commit()
    con.close()


class TestPersonnelR3(unittest.TestCase):
    def test_p5_codes_match_sample_categories(self):
        self.assertEqual(p5_code("Основна щорічна відпустка")[:2], ("В", "08"))
        self.assertEqual(p5_code("Оплачувана тимчасова непрацездатність")[:2], ("ТН", "26"))
        self.assertEqual(p5_code("Відпустка без збереження зарплати за згодою сторін")[:2], ("НА", "18"))
        self.assertIn("Основна щорічна відпустка", ABSENCE_TYPES)

    def test_p5_legend_contains_all_30_codes_from_supplied_sample(self):
        self.assertEqual(len(P5_LEGEND), 30)
        by_code = {code: (label, letter) for label, letter, code in P5_LEGEND}
        self.assertEqual(by_code["10"][1], "Ч")
        self.assertEqual(by_code["20"][1], "НД")
        self.assertEqual(by_code["21"][1], "НП")
        self.assertEqual(by_code["30"][1], "І")

    def test_p5_collector_keeps_weekend_blank_and_counts_absences(self):
        with tempfile.TemporaryDirectory() as tmp:
            db_path = Path(tmp) / "db.sqlite3"
            make_db(db_path)
            core = FakeCore(db_path)
            data = collect_p5_data(core, 2026, 8, active_only=True)
            emp = data["employees"][0]
            self.assertEqual(emp["cells"][2]["code"], "Р")   # 03.08
            self.assertEqual(emp["cells"][3]["code"], "В")   # 04.08
            self.assertEqual(emp["cells"][4]["code"], "ТН")  # 05.08
            self.assertEqual(emp["absence_counts"]["8-10"], 1)
            self.assertEqual(emp["absence_counts"]["26-27"], 1)

    def test_absence_override_sets_zero_plan_without_deleting_source_plan(self):
        with tempfile.TemporaryDirectory() as tmp:
            db_path = Path(tmp) / "db.sqlite3"
            make_db(db_path)
            core = FakeCore(db_path)

            # Replace original day-time with one that would otherwise expose 8h plan.
            def automatic_day(con, employee_id, target_date):
                if target_date == date(2026, 8, 4):
                    return {
                        "day_type": "Робота",
                        "planned_minutes": 480,
                        "actual_minutes": None,
                        "source": "зміна персоналу",
                        "notes": "",
                        "manual": False,
                    }
                return core.employee_day_time(con, employee_id, target_date)

            core.employee_day_time = automatic_day
            # Minimal attributes needed only for class definition.
            core.tk = SimpleNamespace(TclError=Exception)
            core.ttk = SimpleNamespace()
            core.messagebox = SimpleNamespace()
            core.filedialog = SimpleNamespace()
            core.fmt_date = lambda value: value
            core.OUTPUT_DIR = Path(tmp)
            core.write_output_file = lambda *a, **k: None

            Base = type("Base", (), {})
            install(core, Base)
            con = core.db()
            row = core.employee_day_time(con, 1, date(2026, 8, 4))
            con.close()
            self.assertEqual(row["day_type"], "Основна щорічна відпустка")
            self.assertEqual(row["planned_minutes"], 0)
            self.assertEqual(row["actual_minutes"], 0)

    def test_route_schedule_plan_fills_p5_without_question_mark(self):
        with tempfile.TemporaryDirectory() as tmp:
            db_path = Path(tmp) / "db.sqlite3"
            make_db(db_path)
            con = sqlite3.connect(db_path)
            con.execute("UPDATE employees SET driver_id=10 WHERE id=1")
            con.execute("INSERT INTO routes(id,name) VALUES(1,'Маршрут 10')")
            con.execute(
                """INSERT INTO route_segments(
                       id,route_id,segment_no,start_time,end_time,
                       work_start_time,work_end_time,work_hours,driving_hours
                   ) VALUES(1,1,1,'08:00','15:30','07:30','16:15',0,0)"""
            )
            con.execute(
                """INSERT INTO worklog(
                       id,driver_id,work_date,overtime_hours,route_id,
                       work_hours,work_start_time,work_end_time
                   ) VALUES(1,10,'2026-08-06',0,1,0,'','')"""
            )
            con.commit()
            con.close()

            core = FakeCore(db_path)
            core.tk = SimpleNamespace(TclError=Exception)
            core.ttk = SimpleNamespace()
            core.messagebox = SimpleNamespace()
            core.filedialog = SimpleNamespace()
            core.fmt_date = lambda value: value
            core.OUTPUT_DIR = Path(tmp)
            core.write_output_file = lambda *a, **k: None

            Base = type("Base", (), {})
            install(core, Base)

            con = core.db()
            row = core.employee_day_time(con, 1, date(2026, 8, 6))
            con.close()
            self.assertEqual(row["planned_minutes"], 525)
            self.assertIn("графік маршруту", row["source"])

            data = collect_p5_data(core, 2026, 8, active_only=True)
            cell = data["employees"][0]["cells"][5]  # 06.08
            self.assertEqual(cell["code"], "")
            self.assertIsNone(cell["hours"])
            self.assertTrue(cell["missing"])
            self.assertFalse(cell["substituted_plan"])

            substituted = collect_p5_data(
                core, 2026, 8, active_only=True, use_plan_when_fact_missing=True
            )
            subcell = substituted["employees"][0]["cells"][5]
            self.assertEqual(subcell["code"], "Р")
            self.assertEqual(subcell["hours"], 525)
            self.assertFalse(subcell["missing"])
            self.assertTrue(subcell["substituted_plan"])
            self.assertEqual(substituted["planned_substituted_total"], 1)

    def test_duration_only_driver_plan_does_not_invent_clock_time(self):
        with tempfile.TemporaryDirectory() as tmp:
            db_path = Path(tmp) / "db.sqlite3"
            make_db(db_path)
            con = sqlite3.connect(db_path)
            con.row_factory = sqlite3.Row
            con.execute("UPDATE employees SET driver_id=10 WHERE id=1")
            con.execute(
                """INSERT INTO worklog(
                       id,driver_id,work_date,overtime_hours,route_id,
                       work_hours,work_start_time,work_end_time
                   ) VALUES(1,10,'2026-08-06',0,NULL,8,'','')"""
            )
            con.commit()
            employee = con.execute("SELECT * FROM employees WHERE id=1").fetchone()
            core = FakeCore(db_path)
            conflict = driver_plan_conflict(
                core, con, employee,
                datetime(2026,8,6,7,0),
                datetime(2026,8,6,10,0),
            )
            con.close()
            self.assertIsNotNone(conflict)
            self.assertEqual(conflict["kind"], "unknown_time")
            self.assertIsNone(conflict["start"])
            self.assertIsNone(conflict["end"])

    def test_internal_role_allows_non_overlapping_driver_interval(self):
        with tempfile.TemporaryDirectory() as tmp:
            db_path = Path(tmp) / "db.sqlite3"
            make_db(db_path)
            con = sqlite3.connect(db_path)
            con.row_factory = sqlite3.Row
            con.execute("UPDATE employees SET driver_id=10 WHERE id=1")
            con.execute(
                """INSERT INTO worklog(
                       id,driver_id,work_date,overtime_hours,route_id,
                       work_hours,work_start_time,work_end_time
                   ) VALUES(1,10,'2026-08-06',0,NULL,6,'11:00','17:00')"""
            )
            con.commit()
            employee = con.execute("SELECT * FROM employees WHERE id=1").fetchone()
            core = FakeCore(db_path)
            no_conflict = driver_plan_conflict(
                core, con, employee,
                datetime(2026,8,6,7,0),
                datetime(2026,8,6,10,0),
            )
            overlap = driver_plan_conflict(
                core, con, employee,
                datetime(2026,8,6,12,0),
                datetime(2026,8,6,14,0),
            )
            con.close()
            self.assertIsNone(no_conflict)
            self.assertIsNotNone(overlap)
            self.assertEqual(overlap["kind"], "overlap")
            self.assertEqual(overlap["start"].strftime("%H:%M"), "11:00")
            self.assertEqual(overlap["end"].strftime("%H:%M"), "17:00")

    def test_old_driver_balance_reflects_personnel_vacation(self):
        with tempfile.TemporaryDirectory() as tmp:
            db_path = Path(tmp) / "db.sqlite3"
            make_db(db_path)
            con = sqlite3.connect(db_path)
            con.execute("UPDATE employees SET driver_id=10 WHERE id=1")
            con.execute(
                """INSERT INTO worklog(
                       id,driver_id,work_date,overtime_hours,route_id,
                       work_hours,work_start_time,work_end_time
                   ) VALUES(1,10,'2026-08-04',0,NULL,8,'08:00','16:00')"""
            )
            con.commit()
            con.close()

            core = FakeCore(db_path)
            core.tk = SimpleNamespace(TclError=Exception)
            core.ttk = SimpleNamespace()
            core.messagebox = SimpleNamespace()
            core.filedialog = SimpleNamespace()
            core.fmt_date = lambda value: value
            core.OUTPUT_DIR = Path(tmp)
            core.write_output_file = lambda *a, **k: None

            Base = type("Base", (), {})
            install(core, Base)
            data = core.collect_monthly_work_balance(2026, 8, True)
            driver = data["drivers"][0]
            self.assertEqual(driver["cells"][3], "Відп")  # 04.08
            self.assertEqual(driver["work_min"], 30 * 480)
            self.assertEqual(driver["work_days"], 30)
            self.assertEqual(_legacy_absence_cell("Оплачувана тимчасова непрацездатність"), "Лік")

    def test_p5_pdf_and_xlsx_smoke(self):
        with tempfile.TemporaryDirectory() as tmp:
            from openpyxl import load_workbook
            from pypdf import PdfReader

            db_path = Path(tmp) / "db.sqlite3"
            make_db(db_path)
            core = FakeCore(db_path)
            pdf = Path(tmp) / "p5.pdf"
            xlsx = Path(tmp) / "p5.xlsx"
            export_p5_pdf(
                core, 2026, 8, pdf, True,
                date(2026, 8, 31), "Експлуатаційний відділ", "12345678"
            )
            export_p5_xlsx(
                core, 2026, 8, xlsx, True,
                date(2026, 8, 31), "Експлуатаційний відділ", "12345678"
            )
            self.assertTrue(pdf.exists() and pdf.stat().st_size > 1000)
            self.assertTrue(xlsx.exists() and xlsx.stat().st_size > 1000)

            with PdfReader(pdf) as doc:
                self.assertGreaterEqual(len(doc.pages), 2)
                text = "\n".join((page.extract_text() or "") for page in doc.pages)
            # Numeric metadata remains extractable even on CI hosts where the
            # test FakeCore intentionally has no Cyrillic TTF candidate.
            self.assertIn("31.08.2026", text)
            self.assertIn("12345678", text)

            wb = load_workbook(xlsx, data_only=False)
            self.assertIn("Табель П-5", wb.sheetnames)
            self.assertIn("Умовні позначення", wb.sheetnames)
            self.assertIn("ТАБЕЛЬ ОБЛІКУ ВИКОРИСТАННЯ РОБОЧОГО ЧАСУ", str(wb["Табель П-5"]["A2"].value))
            self.assertIn("12345678", str(wb["Табель П-5"]["A3"].value))
            self.assertEqual(wb["Умовні позначення"]["C36"].value, "30")


if __name__ == "__main__":
    unittest.main()
