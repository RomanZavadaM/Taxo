# -*- coding: utf-8 -*-
import sqlite3
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

import military_transport_2026 as mt
import military_transport_statement as statement

ROOT = Path(__file__).resolve().parents[1]


def make_db():
    con = sqlite3.connect(":memory:")
    con.row_factory = sqlite3.Row
    con.executescript("""
        PRAGMA foreign_keys=ON;
        CREATE TABLE vehicles(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            plate TEXT DEFAULT '',
            make_model TEXT DEFAULT '',
            year INTEGER,
            notes TEXT DEFAULT '',
            active INTEGER DEFAULT 1,
            created_at TEXT NOT NULL
        );
        CREATE TABLE employees(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            last_name TEXT NOT NULL,
            first_name TEXT NOT NULL,
            middle_name TEXT DEFAULT '',
            active INTEGER DEFAULT 1
        );
        CREATE TABLE company(
            id INTEGER PRIMARY KEY CHECK(id=1),
            name TEXT DEFAULT '',
            address TEXT DEFAULT '',
            phone TEXT DEFAULT '',
            fax TEXT DEFAULT '',
            email TEXT DEFAULT '',
            signer_name TEXT DEFAULT '',
            signer_position TEXT DEFAULT ''
        );
        INSERT INTO company(id,name,address,signer_name,signer_position)
        VALUES(1,'Тестове АТП','м. Тест','Керівник Тест','Директор');
    """)
    statement.ensure_schema_on_connection(con)
    return con


def add_vehicle(con, *, name="Автобус", plate="AA0001AA", year=2020, active=1):
    cur = con.execute(
        "INSERT INTO vehicles(name,plate,make_model,year,active,created_at) VALUES(?,?,?,?,?,?)",
        (name, plate, "TEST MODEL", year, active, "2026-09-28T00:00:00"),
    )
    return int(cur.lastrowid)


class OwnFleetStatementR1Tests(unittest.TestCase):
    def setUp(self):
        self.con = make_db()

    def tearDown(self):
        self.con.close()

    def test_unknown_scope_is_not_silently_included(self):
        vehicle_id = add_vehicle(self.con)
        overview = statement.statement_overview(self.con, "2026-09-28")
        self.assertEqual(overview[0]["vehicle_id"], vehicle_id)
        self.assertEqual(overview[0]["report_scope"], statement.SCOPE_UNKNOWN)
        self.assertEqual(statement.collect_statement_rows(self.con, "2026-09-28"), [])

    def test_own_scope_includes_vehicle_even_if_operationally_inactive_and_not_designated(self):
        vehicle_id = add_vehicle(self.con, active=0)
        self.assertEqual(mt.profile_state(self.con, vehicle_id)["status"], mt.STATUS_UNKNOWN)
        statement.set_vehicle_statement_data(
            self.con,
            vehicle_id,
            report_scope=statement.SCOPE_OWN,
            vehicle_type="Автобус",
            technical_condition="Справний",
            residual_book_value_thousand_uah="125.5",
        )
        rows = statement.collect_statement_rows(self.con, "2026-09-28")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["plate"], "AA0001AA")
        self.assertEqual(rows[0]["vehicle_type"], "Автобус")
        self.assertEqual(rows[0]["residual_value"], "125.5")

    def test_order_or_designated_status_does_not_put_vehicle_into_own_statement(self):
        vehicle_id = add_vehicle(self.con)
        mt.set_profile(self.con, vehicle_id, status=mt.STATUS_DESIGNATED)
        mt.record_order(self.con, vehicle_id, mt.ORDER_PARTIAL, "2026-09-28", document_no="N-1")
        self.assertEqual(statement.collect_statement_rows(self.con, "2026-09-28"), [])

    def test_explicit_worker_link_populates_personnel_columns(self):
        vehicle_id = add_vehicle(self.con)
        statement.set_vehicle_statement_data(
            self.con, vehicle_id, report_scope=statement.SCOPE_OWN,
            vehicle_type="Автобус", technical_condition="Справний",
            residual_book_value_thousand_uah="100",
        )
        cur = self.con.execute(
            """INSERT INTO employees(last_name,first_name,middle_name,active,birth_date,actual_address)
               VALUES(?,?,?,?,?,?)""",
            ("Тестовий", "Водій", "Один", 1, "1980-04-02", "м. Львів"),
        )
        employee_id = int(cur.lastrowid)
        self.con.execute(
            """INSERT INTO employee_military_profile(
                   employee_id,military_specialty,military_rank
               ) VALUES(?,?,?)""",
            (employee_id, "837037", "солдат"),
        )
        mt.add_worker_link(
            self.con, vehicle_id, employee_id=employee_id,
            relation_note="закріплений водій", valid_from="2026-01-01",
        )
        rows = statement.collect_statement_rows(self.con, date(2026, 9, 28))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["worker_name"], "Тестовий Водій Один")
        self.assertEqual(rows[0]["birth_year"], "1980")
        self.assertEqual(rows[0]["military_specialty"], "837037")
        self.assertEqual(rows[0]["military_rank"], "солдат")
        self.assertEqual(rows[0]["residence"], "м. Львів")

    def test_multiple_workers_repeat_vehicle_as_separate_official_rows(self):
        vehicle_id = add_vehicle(self.con)
        statement.set_vehicle_statement_data(
            self.con, vehicle_id, report_scope=statement.SCOPE_OWN,
            vehicle_type="Автобус", technical_condition="Справний",
            residual_book_value_thousand_uah="100",
        )
        for idx in range(2):
            cur = self.con.execute(
                "INSERT INTO employees(last_name,first_name,middle_name,active) VALUES(?,?,?,1)",
                ("Працівник", str(idx + 1), ""),
            )
            mt.add_worker_link(self.con, vehicle_id, employee_id=int(cur.lastrowid))
        rows = statement.collect_statement_rows(self.con, "2026-09-28")
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["seq"], 1)
        self.assertEqual(rows[1]["seq"], "")

    def test_statement_issues_keep_unknown_and_missing_fields_visible(self):
        add_vehicle(self.con, plate="AA0002AA")
        own = add_vehicle(self.con, plate="AA0003AA")
        statement.set_vehicle_statement_data(self.con, own, report_scope=statement.SCOPE_OWN)
        issues = statement.statement_issues(self.con, "2026-09-28")
        joined = "\n".join(issues)
        self.assertIn("Не визначено належність", joined)
        self.assertIn("AA0003AA", joined)
        self.assertIn("залишкова вартість", joined)

    def test_xlsx_export_contains_appendix_headers(self):
        vehicle_id = add_vehicle(self.con)
        statement.set_vehicle_statement_data(
            self.con, vehicle_id, report_scope=statement.SCOPE_OWN,
            vehicle_type="Автобус", technical_condition="Справний",
            residual_book_value_thousand_uah="100",
        )
        rows = statement.collect_statement_rows(self.con, "2026-09-28")
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "statement.xlsx"
            statement.export_statement_xlsx(
                path, "2026-09-28", statement.company_details(self.con), rows
            )
            self.assertTrue(path.exists())
            wb = load_workbook(path, read_only=True, data_only=True)
            try:
                ws = wb["Відомість"]
                self.assertIn("Відомість про наявність", ws["A1"].value)
                self.assertEqual(ws["A5"].value, "№ з/п")
                self.assertIn("балансова", ws["G5"].value.lower())
            finally:
                wb.close()

    def test_pdf_export_smoke(self):
        vehicle_id = add_vehicle(self.con)
        statement.set_vehicle_statement_data(
            self.con, vehicle_id, report_scope=statement.SCOPE_OWN,
            vehicle_type="Автобус", technical_condition="Справний",
            residual_book_value_thousand_uah="100",
        )
        rows = statement.collect_statement_rows(self.con, "2026-09-28")
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "statement.pdf"
            statement.export_statement_pdf(
                path, "2026-09-28", statement.company_details(self.con), rows
            )
            self.assertTrue(path.exists())
            self.assertGreater(path.stat().st_size, 1000)


class StatementIntegrationR1Tests(unittest.TestCase):
    def test_statement_layer_is_independent_from_shlyakh_and_document_control(self):
        source = (ROOT / "military_transport_statement.py").read_text("utf-8")
        self.assertNotIn("import vehicle_registry", source)
        self.assertNotIn("import vehicle_documents", source)
        self.assertIn("report_scope", source)
        self.assertIn("SCOPE_OWN", source)

    def test_ui_is_wired_into_existing_reports_tab_without_replacing_r10_dashboard(self):
        source = (ROOT / "military_transport_statement_ui.py").read_text("utf-8")
        self.assertIn('anchor = _find_button(win, "Зафіксувати подання")', source)
        self.assertIn('text="Відомість (додаток 1)"', source)
        self.assertIn("Формування документа не означає, що його подано", source)

    def test_feature_layer_remains_in_chain_after_later_revisions(self):
        entry = (ROOT / "taxo_app.py").read_text("utf-8")
        self.assertIn("from v1051_features import install as install_v1051", entry)
        self.assertIn("install_v1051(", entry)
        self.assertIn("install_v10410(", entry)
        layer = (ROOT / "v1051_features.py").read_text("utf-8")
        self.assertIn('APP_VERSION = "10.5-r1"', layer)


if __name__ == "__main__":
    unittest.main()
