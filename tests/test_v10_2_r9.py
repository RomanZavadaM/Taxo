# -*- coding: utf-8 -*-
from pathlib import Path
import tempfile
import unittest

from docx import Document

import main
from release_naming import start_archive_stem, version_from_file


ROOT=Path(__file__).resolve().parents[1]


class TestTaxo102R9AttestationDocx(unittest.TestCase):
    def test_version_is_r9_or_later_in_same_candidate_line(self):
        current=version_from_file(ROOT/"VERSION.txt")
        self.assertEqual(main.APP_VERSION,current)
        self.assertEqual(
            start_archive_stem("10.2-r9"),
            "Taxo_v10_2_candidate_r9_START",
        )

    def test_template_rows_are_found_by_semantic_markers(self):
        doc=Document(str(ROOT/"Бланк підтвердження.docx"))
        ua=doc.tables[0].cell(0,0).paragraphs
        en=doc.tables[1].cell(0,0).paragraphs
        self.assertTrue(main._att_find_paragraph(ua,"12. з ","12. з(").text.strip().startswith("12."))
        self.assertTrue(main._att_find_paragraph(ua,"13. по ","13. по(").text.strip().startswith("13."))
        self.assertTrue(main._att_find_paragraph(ua,"20. Місце ","20.Місце ").text.strip().startswith("20."))
        self.assertTrue(main._att_find_paragraph(en,"12. from ","12.from ").text.strip().startswith("12."))
        self.assertTrue(main._att_find_paragraph(en,"13. to ","13.to ").text.strip().startswith("13."))
        self.assertTrue(main._att_find_paragraph(en,"20. Place ","20.Place ").text.strip().startswith("20."))

    def test_saved_docx_validator_accepts_actual_period_and_date(self):
        doc=Document(str(ROOT/"Бланк підтвердження.docx"))
        ua=doc.tables[0].cell(0,0).paragraphs
        main._att_set_runs(main._att_find_paragraph(ua,"12. з ","12. з("),[
            ("12. з (година/день/місяць/рік): ",False),
            ("20:40 21.09.2026",True),
        ])
        main._att_set_runs(main._att_find_paragraph(ua,"13. по ","13. по("),[
            ("13. по (година/день/місяць/рік): ",False),
            ("07:25 24.09.2026",True),
        ])
        main._att_set_runs(main._att_find_paragraph(ua,"20. Місце ","20.Місце "),[
            ("20. Місце ",False),("с. Муроване",True),("    Дата ",False),("24.09.2026",True),
        ])
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/"attestation.docx"
            doc.save(path)
            main._validate_attestation_docx(
                path,"20:40 21.09.2026","07:25 24.09.2026","24.09.2026"
            )

    def test_validator_rejects_stale_docx_dates(self):
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/"stale.docx"
            Document().save(path)
            with self.assertRaises(RuntimeError):
                main._validate_attestation_docx(
                    path,"20:40 21.09.2026","07:25 24.09.2026","24.09.2026"
                )

    def test_docx_period_fields_do_not_use_hardcoded_paragraph_numbers(self):
        source=(ROOT/"main.py").read_text("utf-8")
        self.assertNotIn("_att_set_runs(up[19]",source)
        self.assertNotIn("_att_set_runs(up[20]",source)
        self.assertNotIn("_att_set_runs(ep[20]",source)
        self.assertNotIn("_att_set_runs(ep[21]",source)
        self.assertIn("_validate_attestation_docx(out_path,period_from,period_to,form_date)",source)


if __name__=="__main__":
    unittest.main()
