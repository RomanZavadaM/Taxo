# -*- coding: utf-8 -*-
"""Taxo 10.4-r8 — import and reconcile vehicles with the state registry «Шлях».

The local Taxo database remains the working source of truth.  The registry is a
periodic external snapshot used to fill missing data, optionally refresh fields,
and show what the state currently knows.  Missing fields/rows in a registry
extract never delete or clear local Taxo data.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from datetime import date, datetime
from pathlib import Path

APP_VERSION = "10.4-r8"
SOURCE_KIND = "shlyakh_vehicle_registry"

IMPORT_COMPARE = "compare"
IMPORT_FILL_EMPTY = "fill_empty"
IMPORT_UPDATE = "update"
IMPORT_MODES = (IMPORT_COMPARE, IMPORT_FILL_EMPTY, IMPORT_UPDATE)
IMPORT_MODE_LABELS = {
    IMPORT_COMPARE: "Лише звірити",
    IMPORT_FILL_EMPTY: "Доповнити порожні поля",
    IMPORT_UPDATE: "Оновити з реєстру + доповнити",
}

FIELD_LABELS = {
    "registry_vehicle_type": "Вид ТЗ",
    "plate": "Державний номер",
    "registry_carrier": "Перевізник",
    "registry_carrier_edrpou": "ЄДРПОУ перевізника",
    "registry_status": "Статус у «Шлях»",
    "vin": "VIN",
    "make": "Марка",
    "model": "Модель",
    "gross_mass_kg": "Повна маса, кг",
    "euro_class": "Екологічний клас EURO",
    "ecmt_euro_class": "EURO за сертифікатом ЄКМТ",
    "ecmt_valid_from": "Початок дії сертифіката ЄКМТ",
    "ecmt_certificate_no": "№ сертифіката ЄКМТ",
}

HEADER_TO_FIELD = {
    "вид": "registry_vehicle_type",
    "№ тз": "plate",
    "номер тз": "plate",
    "перевізник": "registry_carrier",
    "статус": "registry_status",
    "vin код автомобіля": "vin",
    "vin код": "vin",
    "vin": "vin",
    "марка": "make",
    "модель": "model",
    "повна маса кг": "gross_mass_kg",
    "євро": "euro_class",
    "euro згідно сертифікату єкмт": "ecmt_euro_class",
    "початок дії сертифікату єкмт": "ecmt_valid_from",
    "№ сертифікату єкмт": "ecmt_certificate_no",
    "номер сертифікату єкмт": "ecmt_certificate_no",
}

REQUIRED_FIELDS = {"plate", "registry_status", "vin", "make", "model"}


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _norm_header(value):
    text = _text(value).lower().replace("ʼ", "'").replace("’", "'")
    text = text.replace("-", " ").replace("_", " ")
    text = re.sub(r"[^0-9a-zа-яіїєґ№]+", " ", text, flags=re.IGNORECASE)
    return re.sub(r"\s+", " ", text).strip()


_PLATE_TRANSLATE = str.maketrans({
    "А":"A", "В":"B", "С":"C", "Е":"E", "Н":"H", "І":"I",
    "К":"K", "М":"M", "О":"O", "Р":"P", "Т":"T", "Х":"X",
})


def normalize_plate(value):
    text = _text(value).upper().translate(_PLATE_TRANSLATE)
    return re.sub(r"[^A-Z0-9]", "", text)


def normalize_vin(value):
    return re.sub(r"[^A-Z0-9]", "", _text(value).upper())


def _number_text(value):
    text = _text(value).replace(" ", "").replace(",", ".")
    if not text:
        return ""
    try:
        num = float(text)
    except ValueError:
        return _text(value)
    if num.is_integer():
        return str(int(num))
    return ("%.3f" % num).rstrip("0").rstrip(".")


def _iso_date(value):
    if value in (None, ""):
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _text(value)
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return text


def _carrier_edrpou(value):
    match = re.search(r"(?<!\d)(\d{8})(?!\d)", _text(value))
    return match.group(1) if match else ""


def file_sha256(path):
    digest = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _header_mapping(values):
    mapping = {}
    for idx, value in enumerate(values):
        normalized = _norm_header(value)
        field = HEADER_TO_FIELD.get(normalized)
        if field:
            mapping[field] = idx
    return mapping


def _detect_header(rows):
    for idx, row in enumerate(rows[:25]):
        mapping = _header_mapping(row)
        if REQUIRED_FIELDS.issubset(set(mapping)):
            return idx, mapping
    return None, {}


def _value(row, mapping, field):
    idx = mapping.get(field)
    if idx is None or idx >= len(row):
        return ""
    return row[idx]


def _record(row, mapping, row_no):
    data = {
        "registry_vehicle_type": _text(_value(row, mapping, "registry_vehicle_type")),
        "plate": _text(_value(row, mapping, "plate")).upper(),
        "registry_carrier": _text(_value(row, mapping, "registry_carrier")),
        "registry_status": _text(_value(row, mapping, "registry_status")),
        "vin": normalize_vin(_value(row, mapping, "vin")),
        "make": _text(_value(row, mapping, "make")),
        "model": _text(_value(row, mapping, "model")),
        "gross_mass_kg": _number_text(_value(row, mapping, "gross_mass_kg")),
        "euro_class": _text(_value(row, mapping, "euro_class")),
        "ecmt_euro_class": _text(_value(row, mapping, "ecmt_euro_class")),
        "ecmt_valid_from": _iso_date(_value(row, mapping, "ecmt_valid_from")),
        "ecmt_certificate_no": _text(_value(row, mapping, "ecmt_certificate_no")),
    }
    data["registry_carrier_edrpou"] = _carrier_edrpou(data["registry_carrier"])
    data["source_row"] = int(row_no)
    return data


def parse_registry_xlsx(path):
    from openpyxl import load_workbook

    path = Path(path)
    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            header_row, mapping = _detect_header(rows)
            if header_row is None:
                continue
            parsed = []
            for row_no, row in enumerate(rows[header_row + 1:], start=header_row + 2):
                if not any(_text(v) for v in row):
                    continue
                item = _record(row, mapping, row_no)
                if not item["plate"] and not item["vin"]:
                    continue
                parsed.append(item)
            return {
                "source_kind": SOURCE_KIND,
                "source_name": path.name,
                "sheet_name": ws.title,
                "header_row": header_row + 1,
                "rows": parsed,
            }
    finally:
        wb.close()
    raise ValueError("Формат XLSX не розпізнано як експорт транспортних засобів з реєстру «Шлях».")


def _decode_csv(path):
    raw = Path(path).read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            pass
    return raw.decode("utf-8", errors="replace")


def parse_registry_csv(path):
    path = Path(path)
    text = _decode_csv(path)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";"
    rows = list(csv.reader(io.StringIO(text), dialect))
    header_row, mapping = _detect_header(rows)
    if header_row is None:
        raise ValueError("Формат CSV не розпізнано як експорт транспортних засобів з реєстру «Шлях».")
    parsed = []
    for row_no, row in enumerate(rows[header_row + 1:], start=header_row + 2):
        if not any(_text(v) for v in row):
            continue
        item = _record(row, mapping, row_no)
        if not item["plate"] and not item["vin"]:
            continue
        parsed.append(item)
    return {
        "source_kind": SOURCE_KIND,
        "source_name": path.name,
        "sheet_name": "CSV",
        "header_row": header_row + 1,
        "rows": parsed,
    }


def parse_registry_file(path):
    suffix = Path(path).suffix.lower()
    if suffix == ".xlsx":
        return parse_registry_xlsx(path)
    if suffix == ".csv":
        return parse_registry_csv(path)
    raise ValueError("Підтримуються файли XLSX або CSV з реєстру «Шлях».")


def ensure_schema_on_connection(con):
    columns = {row[1] for row in con.execute("PRAGMA table_info(vehicles)").fetchall()}
    additions = (
        ("vin", "TEXT DEFAULT ''"),
        ("make", "TEXT DEFAULT ''"),
        ("model", "TEXT DEFAULT ''"),
        ("gross_mass_kg", "TEXT DEFAULT ''"),
        ("euro_class", "TEXT DEFAULT ''"),
        ("registry_vehicle_type", "TEXT DEFAULT ''"),
        ("registry_carrier", "TEXT DEFAULT ''"),
        ("registry_carrier_edrpou", "TEXT DEFAULT ''"),
        ("registry_status", "TEXT DEFAULT ''"),
        ("ecmt_euro_class", "TEXT DEFAULT ''"),
        ("ecmt_valid_from", "TEXT DEFAULT ''"),
        ("ecmt_certificate_no", "TEXT DEFAULT ''"),
        ("registry_last_source_name", "TEXT DEFAULT ''"),
        ("registry_last_verified_at", "TEXT DEFAULT ''"),
    )
    for name, ddl in additions:
        if name not in columns:
            con.execute("ALTER TABLE vehicles ADD COLUMN %s %s" % (name, ddl))
    con.execute("CREATE INDEX IF NOT EXISTS idx_vehicles_vin ON vehicles(vin)")
    con.executescript(
        """
        CREATE TABLE IF NOT EXISTS vehicle_registry_imports (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            source_kind TEXT NOT NULL,
            source_name TEXT NOT NULL,
            file_sha256 TEXT NOT NULL,
            imported_at TEXT NOT NULL,
            mode TEXT NOT NULL,
            total_rows INTEGER NOT NULL DEFAULT 0,
            matched_rows INTEGER NOT NULL DEFAULT 0,
            created_rows INTEGER NOT NULL DEFAULT 0,
            updated_rows INTEGER NOT NULL DEFAULT 0,
            skipped_rows INTEGER NOT NULL DEFAULT 0,
            conflict_rows INTEGER NOT NULL DEFAULT 0
        );
        CREATE INDEX IF NOT EXISTS idx_vehicle_registry_imports_date
            ON vehicle_registry_imports(imported_at DESC);
        CREATE TABLE IF NOT EXISTS vehicle_registry_history (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            import_id INTEGER REFERENCES vehicle_registry_imports(id) ON DELETE SET NULL,
            vehicle_id INTEGER REFERENCES vehicles(id) ON DELETE SET NULL,
            source_kind TEXT NOT NULL,
            source_name TEXT NOT NULL,
            file_sha256 TEXT NOT NULL,
            source_row INTEGER,
            row_key TEXT NOT NULL,
            snapshot_json TEXT NOT NULL,
            imported_at TEXT NOT NULL
        );
        CREATE INDEX IF NOT EXISTS idx_vehicle_registry_history_vehicle
            ON vehicle_registry_history(vehicle_id,imported_at DESC);
        """
    )


def ensure_schema(core):
    con = core.db()
    try:
        ensure_schema_on_connection(con)
        con.commit()
    finally:
        con.close()


def _row_dict(row):
    return {key: row[key] for key in row.keys()}


def _norm_value(field, value):
    if field == "plate":
        return normalize_plate(value)
    if field == "vin":
        return normalize_vin(value)
    if field == "gross_mass_kg":
        return _number_text(value)
    return _text(value).casefold()


def _vehicle_rows(con):
    return con.execute("SELECT * FROM vehicles ORDER BY active DESC,id").fetchall()


def _match_vehicle(con, item):
    vehicles = _vehicle_rows(con)
    vin = normalize_vin(item.get("vin"))
    plate = normalize_plate(item.get("plate"))
    vin_matches = [row for row in vehicles if vin and normalize_vin(row["vin"] if "vin" in row.keys() else "") == vin]
    plate_matches = [row for row in vehicles if plate and normalize_plate(row["plate"]) == plate]

    if len(vin_matches) > 1:
        return None, "critical", "У Taxo кілька ТЗ мають однаковий VIN"
    if len(plate_matches) > 1:
        return None, "critical", "У Taxo кілька ТЗ мають однаковий державний номер"

    vin_row = vin_matches[0] if vin_matches else None
    plate_row = plate_matches[0] if plate_matches else None
    if vin_row is not None and plate_row is not None and int(vin_row["id"]) != int(plate_row["id"]):
        return None, "critical", "VIN і державний номер вказують на різні картки Taxo"
    if vin_row is not None:
        return vin_row, "vin", ""
    if plate_row is not None:
        local_vin = normalize_vin(plate_row["vin"] if "vin" in plate_row.keys() else "")
        if vin and local_vin and vin != local_vin:
            return None, "critical", "Державний номер збігається, але VIN відрізняється"
        return plate_row, "plate", ""
    return None, "new_registry", "ТЗ є у реєстрі, але картки в Taxo ще немає"


def _change(field, old, new):
    return {
        "field": field,
        "label": FIELD_LABELS.get(field, field),
        "old": _text(old),
        "new": _text(new),
        "kind": "fill" if not _text(old) and _text(new) else "difference",
    }


def _incoming_fields(item):
    return {field: _text(item.get(field)) for field in FIELD_LABELS if _text(item.get(field))}


def plan_import_rows(con, parsed):
    ensure_schema_on_connection(con)
    plan = []
    for item in parsed["rows"]:
        vehicle, match_quality, note = _match_vehicle(con, item)
        if vehicle is None:
            status = "critical" if match_quality == "critical" else "new_registry"
            plan.append({
                "item": item,
                "vehicle_id": None,
                "vehicle_name": "%s %s — %s" % (item.get("make", ""), item.get("model", ""), item.get("plate", "")),
                "match_quality": match_quality,
                "status": status,
                "changes": [],
                "notes": [note] if note else [],
            })
            continue

        current = _row_dict(vehicle)
        changes = []
        for field, new in _incoming_fields(item).items():
            old = current.get(field, "")
            if _norm_value(field, old) != _norm_value(field, new):
                changes.append(_change(field, old, new))
        if not changes:
            status = "no_changes"
        elif all(change["kind"] == "fill" for change in changes):
            status = "fill"
        else:
            status = "difference"
        plan.append({
            "item": item,
            "vehicle_id": int(vehicle["id"]),
            "vehicle_name": " — ".join(x for x in (vehicle["name"], vehicle["plate"], vehicle["make_model"]) if _text(x)),
            "match_quality": match_quality,
            "status": status,
            "changes": changes,
            "notes": [],
        })
    return plan


def preview_registry_import(core, path):
    parsed = parse_registry_file(path)
    sha = file_sha256(path)
    con = core.db()
    try:
        plan = plan_import_rows(con, parsed)
        matched_ids = {int(row["vehicle_id"]) for row in plan if row.get("vehicle_id") is not None}
        local_only = []
        for vehicle in _vehicle_rows(con):
            if int(vehicle["id"]) in matched_ids:
                continue
            local_only.append({
                "vehicle_id": int(vehicle["id"]),
                "vehicle_name": " — ".join(x for x in (vehicle["name"], vehicle["plate"], vehicle["make_model"]) if _text(x)),
                "status": "local_only",
                "note": "Є у Taxo, але відсутнє у цьому витягу. Дані не видаляються і не обнуляються.",
            })
    finally:
        con.close()
    return {"path": str(path), "file_sha256": sha, "parsed": parsed, "plan": plan, "local_only": local_only}


def _changes_for_mode(changes, mode):
    if mode not in IMPORT_MODES:
        raise ValueError("Невідомий режим реєстрової звірки: %s" % mode)
    if mode == IMPORT_COMPARE:
        return []
    if mode == IMPORT_FILL_EMPTY:
        return [change for change in changes if change["kind"] == "fill" and _text(change["new"])]
    return [change for change in changes if _text(change["new"])]


def _row_key(item):
    vin = normalize_vin(item.get("vin"))
    plate = normalize_plate(item.get("plate"))
    return vin or plate or "row-%s" % item.get("source_row", "")


def _snapshot(item):
    return {key: _text(value) for key, value in sorted(item.items()) if key != "source_row" and _text(value)}


def _new_vehicle_name(item):
    label = " ".join(x for x in (_text(item.get("make")), _text(item.get("model"))) if x).strip()
    return label or _text(item.get("plate")) or _text(item.get("vin")) or "Транспортний засіб"


def _create_vehicle(con, item, source_name, now):
    make_model = " ".join(x for x in (_text(item.get("make")), _text(item.get("model"))) if x).strip()
    values = {
        "name": _new_vehicle_name(item),
        "plate": _text(item.get("plate")),
        "make_model": make_model,
        "active": 1,
        "created_at": now,
    }
    for field in FIELD_LABELS:
        if field == "plate":
            continue
        values[field] = _text(item.get(field))
    values["registry_last_source_name"] = source_name
    values["registry_last_verified_at"] = now
    columns = list(values)
    placeholders = ",".join("?" for _ in columns)
    cur = con.execute(
        "INSERT INTO vehicles(%s) VALUES(%s)" % (",".join(columns), placeholders),
        [values[column] for column in columns],
    )
    return int(cur.lastrowid)


def apply_registry_import(core, preview, mode=IMPORT_FILL_EMPTY, create_new=False):
    if mode not in IMPORT_MODES:
        raise ValueError("Невідомий режим реєстрової звірки: %s" % mode)
    path = preview["path"]
    current_sha = file_sha256(path)
    if current_sha != preview["file_sha256"]:
        raise ValueError("Файл змінився після попереднього перегляду. Виконайте звірку ще раз.")
    parsed = preview["parsed"]
    plan = preview["plan"]
    now = datetime.now().isoformat(timespec="seconds")
    con = core.db()
    try:
        ensure_schema_on_connection(con)
        cur = con.execute(
            """INSERT INTO vehicle_registry_imports(
                   source_kind,source_name,file_sha256,imported_at,mode,total_rows
               ) VALUES(?,?,?,?,?,?)""",
            (SOURCE_KIND, parsed["source_name"], current_sha, now, mode, len(plan)),
        )
        import_id = int(cur.lastrowid)
        matched = created = updated = skipped = conflicts = 0

        for row_plan in plan:
            item = row_plan["item"]
            vehicle_id = row_plan.get("vehicle_id")
            status = row_plan["status"]
            if status == "critical":
                conflicts += 1
                skipped += 1
            elif vehicle_id is None:
                if create_new and mode != IMPORT_COMPARE and status == "new_registry":
                    vehicle_id = _create_vehicle(con, item, parsed["source_name"], now)
                    created += 1
                else:
                    skipped += 1
            else:
                matched += 1
                selected = _changes_for_mode(row_plan.get("changes") or [], mode)
                if selected:
                    sql = "UPDATE vehicles SET " + ", ".join(change["field"] + "=?" for change in selected)
                    sql += ", registry_last_source_name=?, registry_last_verified_at=? WHERE id=?"
                    con.execute(
                        sql,
                        [change["new"] for change in selected] + [parsed["source_name"], now, vehicle_id],
                    )
                    updated += 1
                elif mode != IMPORT_COMPARE:
                    con.execute(
                        "UPDATE vehicles SET registry_last_source_name=?,registry_last_verified_at=? WHERE id=?",
                        (parsed["source_name"], now, vehicle_id),
                    )

            con.execute(
                """INSERT INTO vehicle_registry_history(
                       import_id,vehicle_id,source_kind,source_name,file_sha256,source_row,row_key,
                       snapshot_json,imported_at
                   ) VALUES(?,?,?,?,?,?,?,?,?)""",
                (
                    import_id, vehicle_id, SOURCE_KIND, parsed["source_name"], current_sha,
                    item.get("source_row"), _row_key(item),
                    json.dumps(_snapshot(item), ensure_ascii=False, sort_keys=True), now,
                ),
            )

        con.execute(
            """UPDATE vehicle_registry_imports
               SET matched_rows=?,created_rows=?,updated_rows=?,skipped_rows=?,conflict_rows=?
               WHERE id=?""",
            (matched, created, updated, skipped, conflicts, import_id),
        )
        con.commit()
        return {
            "import_id": import_id,
            "total": len(plan),
            "matched": matched,
            "created": created,
            "updated": updated,
            "skipped": skipped,
            "conflicts": conflicts,
            "local_only": len(preview.get("local_only") or []),
            "mode": mode,
            "source_name": parsed["source_name"],
        }
    except Exception:
        con.rollback()
        raise
    finally:
        con.close()


def registry_quarter_status(con, today=None):
    ensure_schema_on_connection(con)
    today = today or date.today()
    quarter = (today.month - 1) // 3 + 1
    first_month = (quarter - 1) * 3 + 1
    quarter_start = date(today.year, first_month, 1)
    row = con.execute(
        "SELECT * FROM vehicle_registry_imports ORDER BY imported_at DESC,id DESC LIMIT 1"
    ).fetchone()
    last_date = None
    if row is not None:
        try:
            last_date = datetime.fromisoformat(str(row["imported_at"])).date()
        except Exception:
            pass
    current = bool(last_date and quarter_start <= last_date <= today)
    return {
        "current": current,
        "quarter": "Q%d %d" % (quarter, today.year),
        "label": "Актуально" if current else "Потрібне звіряння",
        "last_date": last_date.isoformat() if last_date else "",
        "source_name": row["source_name"] if row is not None else "",
        "mode": row["mode"] if row is not None else "",
    }


def vehicle_registry_data(con, vehicle_id):
    ensure_schema_on_connection(con)
    row = con.execute("SELECT * FROM vehicles WHERE id=?", (int(vehicle_id),)).fetchone()
    if row is None:
        return None
    result = {field: row[field] if field in row.keys() else "" for field in FIELD_LABELS}
    result["registry_last_source_name"] = row["registry_last_source_name"] if "registry_last_source_name" in row.keys() else ""
    result["registry_last_verified_at"] = row["registry_last_verified_at"] if "registry_last_verified_at" in row.keys() else ""
    return result
