# -*- coding: utf-8 -*-
"""Вбудований переглядач документів Taxo."""
from __future__ import annotations

import os
import shutil
import subprocess
import sys
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

import fitz
from PIL import Image, ImageTk
from docx import Document


PDF_EXTENSIONS = {".pdf"}
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp", ".webp"}
DOCX_EXTENSIONS = {".docx"}
PREVIEW_EXTENSIONS = PDF_EXTENSIONS | IMAGE_EXTENSIONS | DOCX_EXTENSIONS


def document_kind(path):
    suffix = Path(path).suffix.lower()
    if suffix in PDF_EXTENSIONS:
        return "pdf"
    if suffix in IMAGE_EXTENSIONS:
        return "image"
    if suffix in DOCX_EXTENSIONS:
        return "docx"
    return "other"


def system_open(path):
    target = str(Path(path))
    if os.name == "nt":
        os.startfile(target)
    elif sys.platform == "darwin":
        subprocess.Popen(["open", target])
    else:
        subprocess.Popen(["xdg-open", target])


def _windows_print_raster(path):
    import win32con
    import win32print
    import win32ui
    from PIL import ImageWin

    target = Path(path)
    printer_name = win32print.GetDefaultPrinter()
    dc = win32ui.CreateDC()
    dc.CreatePrinterDC(printer_name)
    printable_w = dc.GetDeviceCaps(win32con.HORZRES)
    printable_h = dc.GetDeviceCaps(win32con.VERTRES)

    def draw_image(image):
        image = image.convert("RGB")
        scale = min(printable_w / image.width, printable_h / image.height)
        width = max(1, int(image.width * scale))
        height = max(1, int(image.height * scale))
        left = max(0, (printable_w - width) // 2)
        top = max(0, (printable_h - height) // 2)
        dib = ImageWin.Dib(image)
        dc.StartPage()
        dib.draw(dc.GetHandleOutput(), (left, top, left + width, top + height))
        dc.EndPage()

    dc.StartDoc(target.name)
    try:
        if document_kind(target) == "pdf":
            pdf = fitz.open(str(target))
            try:
                for page in pdf:
                    pix = page.get_pixmap(matrix=fitz.Matrix(2.5, 2.5), alpha=False)
                    image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                    draw_image(image)
            finally:
                pdf.close()
        else:
            with Image.open(target) as image:
                draw_image(image.copy())
    finally:
        dc.EndDoc()
        dc.DeleteDC()


def system_print(path):
    target = str(Path(path))
    if os.name == "nt":
        _windows_print_raster(target)
    else:
        subprocess.Popen(["lp", target])


def docx_preview_text(path):
    """Змістовний preview DOCX; не точне відтворення сторінки Word."""
    doc = Document(str(path))
    chunks = []
    for paragraph in doc.paragraphs:
        value = paragraph.text.rstrip()
        if value:
            chunks.append(value)
    for table in doc.tables:
        if chunks:
            chunks.append("")
        for row in table.rows:
            chunks.append(" | ".join(
                cell.text.strip().replace("\n", " ") for cell in row.cells
            ))
    return "\n".join(chunks).strip()


def _save_copy(parent, path):
    source = Path(path)
    target = filedialog.asksaveasfilename(
        parent=parent,
        title="Зберегти копію документа",
        initialfile=source.name,
        defaultextension=source.suffix,
        filetypes=[("Документ", f"*{source.suffix}"), ("Усі файли", "*.*")],
    )
    if not target:
        return
    try:
        shutil.copy2(source, target)
    except OSError as exc:
        messagebox.showerror("Зберегти копію", str(exc), parent=parent)


def _external(parent, path, external_opener=None):
    try:
        (external_opener or system_open)(path)
    except Exception as exc:
        messagebox.showerror(
            "Відкрити зовнішньо",
            f"Не вдалося відкрити документ:\n{exc}",
            parent=parent,
        )


def _print(parent, path):
    try:
        system_print(path)
        messagebox.showinfo("Друк", "Документ передано системі друку.", parent=parent)
    except Exception as exc:
        messagebox.showwarning(
            "Друк",
            "Не вдалося передати документ системі друку. "
            "Скористайтеся «Відкрити зовнішньо».\n\n" + str(exc),
            parent=parent,
        )


class RasterDocumentWindow:
    def __init__(self, parent, path, external_opener=None):
        self.path = Path(path)
        self.external_opener = external_opener
        self.kind = document_kind(self.path)
        self.page_index = 0
        self.zoom = 1.15
        self.fit_width = True
        self.photo = None
        self.pdf = fitz.open(str(self.path)) if self.kind == "pdf" else None
        self.page_count = len(self.pdf) if self.pdf is not None else 1
        self.image = None if self.kind == "pdf" else Image.open(self.path)

        self.win = tk.Toplevel(parent)
        self.win.title(f"Taxo — перегляд: {self.path.name}")
        self.win.geometry("1050x780")
        self.win.minsize(720, 520)
        if parent is not None:
            try:
                self.win.transient(parent)
            except tk.TclError:
                pass
        self.win.protocol("WM_DELETE_WINDOW", self.close)

        toolbar = ttk.Frame(self.win, padding=(8, 7))
        toolbar.pack(fill="x")
        self.prev_button = ttk.Button(toolbar, text="◀", width=4, command=self.prev_page)
        self.prev_button.pack(side="left")
        self.next_button = ttk.Button(toolbar, text="▶", width=4, command=self.next_page)
        self.next_button.pack(side="left", padx=(4, 8))
        self.page_var = tk.StringVar()
        ttk.Label(toolbar, textvariable=self.page_var, width=18).pack(side="left")
        ttk.Button(toolbar, text="−", width=4, command=lambda: self.change_zoom(-0.15)).pack(side="left")
        ttk.Button(toolbar, text="+", width=4, command=lambda: self.change_zoom(0.15)).pack(side="left", padx=(4, 4))
        ttk.Button(toolbar, text="По ширині", command=self.set_fit_width).pack(side="left", padx=(4, 12))
        self.zoom_var = tk.StringVar()
        ttk.Label(toolbar, textvariable=self.zoom_var, width=9).pack(side="left")
        ttk.Button(toolbar, text="Друк", command=lambda: _print(self.win, self.path)).pack(side="right", padx=(4, 0))
        ttk.Button(
            toolbar, text="Відкрити зовнішньо",
            command=lambda: _external(self.win, self.path, self.external_opener),
        ).pack(side="right", padx=4)
        ttk.Button(
            toolbar, text="Зберегти копію",
            command=lambda: _save_copy(self.win, self.path),
        ).pack(side="right", padx=4)

        body = ttk.Frame(self.win)
        body.pack(fill="both", expand=True)
        self.canvas = tk.Canvas(body, background="#6C7075", highlightthickness=0)
        ybar = ttk.Scrollbar(body, orient="vertical", command=self.canvas.yview)
        xbar = ttk.Scrollbar(body, orient="horizontal", command=self.canvas.xview)
        self.canvas.configure(yscrollcommand=ybar.set, xscrollcommand=xbar.set)
        self.canvas.grid(row=0, column=0, sticky="nsew")
        ybar.grid(row=0, column=1, sticky="ns")
        xbar.grid(row=1, column=0, sticky="ew")
        body.rowconfigure(0, weight=1)
        body.columnconfigure(0, weight=1)
        self.canvas.bind("<Configure>", self._on_resize)
        self.win.bind("<Left>", lambda _e: self.prev_page())
        self.win.bind("<Right>", lambda _e: self.next_page())
        self.render()

    def close(self):
        try:
            if self.pdf is not None:
                self.pdf.close()
            if self.image is not None:
                self.image.close()
        finally:
            self.win.destroy()

    def _on_resize(self, _event=None):
        if self.fit_width:
            self.win.after_idle(self.render)

    def _source_size(self):
        if self.kind == "pdf":
            rect = self.pdf[self.page_index].rect
            return float(rect.width), float(rect.height)
        return float(self.image.width), float(self.image.height)

    def _effective_zoom(self):
        if not self.fit_width:
            return self.zoom
        source_width, _ = self._source_size()
        canvas_width = max(220, self.canvas.winfo_width() - 36)
        return max(0.2, min(4.0, canvas_width / source_width))

    def render(self):
        scale = self._effective_zoom()
        if self.kind == "pdf":
            page = self.pdf[self.page_index]
            pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
            image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        else:
            size = (
                max(1, int(self.image.width * scale)),
                max(1, int(self.image.height * scale)),
            )
            image = self.image.resize(size, Image.Resampling.LANCZOS)
        self.photo = ImageTk.PhotoImage(image)
        self.canvas.delete("all")
        self.canvas.create_image(18, 18, image=self.photo, anchor="nw")
        self.canvas.configure(
            scrollregion=(0, 0, self.photo.width() + 36, self.photo.height() + 36)
        )
        self.page_var.set(
            f"Сторінка {self.page_index + 1} / {self.page_count}"
            if self.kind == "pdf" else "Зображення"
        )
        self.zoom_var.set(f"{int(scale * 100)}%")
        self.prev_button.configure(state=("normal" if self.page_index > 0 else "disabled"))
        self.next_button.configure(
            state=("normal" if self.page_index + 1 < self.page_count else "disabled")
        )

    def prev_page(self):
        if self.page_index > 0:
            self.page_index -= 1
            self.render()

    def next_page(self):
        if self.page_index + 1 < self.page_count:
            self.page_index += 1
            self.render()

    def change_zoom(self, delta):
        self.fit_width = False
        self.zoom = max(0.25, min(4.0, self.zoom + delta))
        self.render()

    def set_fit_width(self):
        self.fit_width = True
        self.render()


class DocxDocumentWindow:
    def __init__(self, parent, path, external_opener=None):
        self.path = Path(path)
        self.external_opener = external_opener
        self.win = tk.Toplevel(parent)
        self.win.title(f"Taxo — DOCX: {self.path.name}")
        self.win.geometry("980x720")
        self.win.minsize(700, 500)
        if parent is not None:
            try:
                self.win.transient(parent)
            except tk.TclError:
                pass

        toolbar = ttk.Frame(self.win, padding=(8, 7))
        toolbar.pack(fill="x")
        ttk.Label(
            toolbar,
            text="Змістовний перегляд DOCX — макет може відрізнятися від Word/PDF.",
        ).pack(side="left")
        ttk.Button(
            toolbar, text="Відкрити оригінал",
            command=lambda: _external(self.win, self.path, self.external_opener),
        ).pack(side="right")
        ttk.Button(
            toolbar, text="Зберегти копію",
            command=lambda: _save_copy(self.win, self.path),
        ).pack(side="right", padx=4)

        frame = ttk.Frame(self.win)
        frame.pack(fill="both", expand=True, padx=8, pady=(0, 8))
        text = tk.Text(frame, wrap="word", padx=18, pady=14, undo=False)
        ybar = ttk.Scrollbar(frame, orient="vertical", command=text.yview)
        text.configure(yscrollcommand=ybar.set)
        text.pack(side="left", fill="both", expand=True)
        ybar.pack(side="right", fill="y")
        try:
            preview = docx_preview_text(self.path)
        except Exception as exc:
            preview = f"Не вдалося прочитати DOCX:\n{exc}"
        text.insert("1.0", preview or "(Документ не містить тексту для preview.)")
        text.configure(state="disabled")


def open_document(parent, path, external_opener=None):
    target = Path(path)
    opener = external_opener or system_open
    # Non-GUI callers/tests and early-startup contexts must never crash while
    # trying to construct a Toplevel without a valid Tk parent. Delegate to the
    # supplied opener before GUI-specific file validation.
    if parent is not None and not hasattr(parent, "tk"):
        return opener(target)
    if not target.exists():
        raise FileNotFoundError(str(target))
    kind = document_kind(target)
    if kind in {"pdf", "image"}:
        return RasterDocumentWindow(parent, target, external_opener=external_opener).win
    if kind == "docx":
        return DocxDocumentWindow(parent, target, external_opener=external_opener).win
    return opener(target)
