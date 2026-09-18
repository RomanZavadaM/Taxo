# -*- coding: utf-8 -*-
"""Personnel P-5 style reporting for Taxo 9.1.

The layout and symbolic codes follow the user-provided August 2026 P-5 sample.
The module does not alter the database schema; it reads employees, daily
timesheet entries and existing work/shift plans.
"""
from __future__ import annotations

import calendar
import os
from datetime import date, datetime
from pathlib import Path
from xml.sax.saxutils import escape


P5_CODES = [
    ("Години роботи, передбачені колдоговором", "Р", "01"),
    ("Години роботи працівників, яким встановлено неповний робочий день (тиждень) згідно з законодавством", "РС", "02"),
    ("Вечірні години роботи", "ВЧ", "03"),
    ("Нічні години роботи", "РН", "04"),
    ("Надурочні години роботи", "НУ", "05"),
    ("Години роботи у вихідні та святкові дні", "РВ", "06"),
    ("Відрядження", "ВД", "07"),
    ("Основна щорічна відпустка", "В", "08"),
    ("Щорічна додаткова відпустка", "Д", "09"),
    ("Додаткова відпустка постраждалим внаслідок Чорнобильської катастрофи", "Ч", "10"),
    ("Творча відпустка", "ТВ", "11"),
    ("Додаткова відпустка у зв'язку з навчанням", "Н", "12"),
    ("Відпустка без збереження заробітної плати у зв'язку з навчанням", "НБ", "13"),
    ("Додаткова відпустка без збереження заробітної плати в обов'язковому порядку", "ДБ", "14"),
    ("Додаткова оплачувана відпустка працівникам, які мають дітей", "ДО", "15"),
    ("Відпустка у зв'язку з вагітністю і пологами / догляд за дитиною до 3 років", "ВП", "16"),
    ("Відпустка для догляду за дитиною до досягнення нею 6-річного віку", "ДД", "17"),
    ("Відпустка без збереження заробітної плати за згодою сторін", "НА", "18"),
    ("Інші відпустки без збереження заробітної плати (на період припинення виконання робіт)", "БЗ", "19"),
    ("Неявки у зв'язку з переведенням за ініціативою роботодавця на неповний робочий день (тиждень)", "НД", "20"),
    ("Неявки у зв'язку з тимчасовим переведенням на роботу на інше підприємство", "НП", "21"),
    ("Інший невідпрацьований час, передбачений законодавством", "ІН", "22"),
    ("Простої", "П", "23"),
    ("Прогули", "ПР", "24"),
    ("Масові невиходи на роботу (страйки)", "С", "25"),
    ("Оплачувана тимчасова непрацездатність", "ТН", "26"),
    ("Неоплачувана тимчасова непрацездатність", "НН", "27"),
    ("Неявки з нез'ясованих причин", "НЗ", "28"),
    ("Інші види неявок, передбачені колективними договорами, угодами", "ІВ", "29"),
    ("Інші причини неявок", "І", "30"),
]

P5_BY_LABEL = {label: (letter, code) for label, letter, code in P5_CODES}
P5_BY_CODE = {code: (label, letter) for label, letter, code in P5_CODES}
P5_BY_LETTER = {letter: (label, code) for label, letter, code in P5_CODES}

LEGACY_DAY_TYPE_ALIASES = {
    "Відпустка": "Основна щорічна відпустка",
    "Лікарняний": "Оплачувана тимчасова непрацездатність",
}

ABSENCE_CODES = {f"{n:02d}" for n in range(8, 31)}
ABSENCE_TYPES = {
    label for label, _letter, code in P5_CODES if code in ABSENCE_CODES
}
ABSENCE_TYPES.update({"Відпустка", "Лікарняний"})

COMMON_ABSENCE_TYPES = [
    "Основна щорічна відпустка",
    "Відпустка без збереження заробітної плати за згодою сторін",
    "Оплачувана тимчасова непрацездатність",
    "Неоплачувана тимчасова непрацездатність",
    "Інші причини неявок",
]


def canonical_day_type(value: str) -> str:
    raw = str(value or "").strip()
    return LEGACY_DAY_TYPE_ALIASES.get(raw, raw)


def is_absence_day_type(value: str) -> bool:
    return canonical_day_type(value) in {
        canonical_day_type(item) for item in ABSENCE_TYPES
    }


def p5_code_for_day_type(value: str):
    canonical = canonical_day_type(value)
    return P5_BY_LABEL.get(canonical)


def format_hours_minutes(minutes):
    if minutes is None:
        return ""
    minutes = int(round(minutes))
    if minutes <= 0:
        return ""
    hours, mins = divmod(minutes, 60)
    return str(hours) if mins == 0 else f"{hours}:{mins:02d}"


def employee_automatic_plan_minutes(core, con, employee, target_date):
    driver_minutes = core._driver_plan_minutes_for_day(
        con, employee["driver_id"], target_date
    )
    shift_minutes, _shift_actual, _found = core._employee_shift_minutes_for_day(
        con, employee["id"], target_date
    )
    return int(driver_minutes + shift_minutes)


def _absence_bucket(code):
    n = int(code)
    if 8 <= n <= 10:
        return "vacation"
    if n in {11, 12, 13, 14, 15, 17, 22}:
        return "other_leave"
    if n == 18:
        return "unpaid_agreement"
    if n == 19:
        return "unpaid_shutdown"
    if n == 20:
        return "part_time_transfer"
    if n == 21:
        return "other_company"
    if n == 23:
        return "downtime"
    if n == 24:
        return "absence"
    if n == 25:
        return "strike"
    if n in {26, 27}:
        return "sick"
    if n in {28, 29, 30}:
        return "other_absence"
    return None


P5_ABSENCE_BUCKETS = [
    ("vacation", "Відпустки\n08-10"),
    ("other_leave", "Інші відпустки\n11-15,17,22"),
    ("unpaid_agreement", "Без з/п\n18"),
    ("unpaid_shutdown", "Без з/п\n19"),
    ("part_time_transfer", "Неповн. день\n20"),
    ("other_company", "Інше підпр.\n21"),
    ("downtime", "Простій\n23"),
    ("absence", "Прогули\n24"),
    ("strike", "Страйки\n25"),
    ("sick", "Непрацезд.\n26-27"),
    ("other_absence", "Інші\n28-30"),
]


def collect_p5(core, year, month, active_only=True):
    y, m = int(year), int(month)
    days = [date(y, m, d) for d in range(1, calendar.monthrange(y, m)[1] + 1)]
    con = core.db()
    sql = """SELECT e.*,GROUP_CONCAT(er.role, ', ') roles
             FROM employees e
             LEFT JOIN employee_roles er ON er.employee_id=e.id"""
    if active_only:
        sql += " WHERE e.active=1"
    sql += " GROUP BY e.id ORDER BY e.active DESC,e.last_name,e.first_name,e.middle_name"
    employees = con.execute(sql).fetchall()
    company = con.execute("SELECT * FROM company WHERE id=1").fetchone()

    rows = []
    for employee in employees:
        if not any(core.employee_employed_on(employee, d) for d in days):
            continue
        cells = []
        work_days = 0
        work_minutes = 0
        overtime_minutes = 0
        weekend_minutes = 0
        absence = {
            key: {"days": 0, "minutes": 0}
            for key, _label in P5_ABSENCE_BUCKETS
        }
        missing_fact = 0

        for d in days:
            if not core.employee_employed_on(employee, d):
                cells.append({"mark": "", "hours": "", "type": "—"})
                continue

            row = core.employee_day_time(con, employee["id"], d)
            day_type = canonical_day_type(row["day_type"])
            code_info = p5_code_for_day_type(day_type)
            auto_plan = employee_automatic_plan_minutes(core, con, employee, d)
            actual = row["actual_minutes"]

            if code_info and code_info[1] in ABSENCE_CODES:
                letter, code = code_info
                cells.append({"mark": letter, "hours": "", "type": day_type})
                bucket = _absence_bucket(code)
                if bucket:
                    absence[bucket]["days"] += 1
                    absence[bucket]["minutes"] += auto_plan
                continue

            if code_info and code_info[1] == "07":
                letter, _code = code_info
                shown_minutes = actual if actual is not None else row["planned_minutes"]
                cells.append({
                    "mark": letter,
                    "hours": format_hours_minutes(shown_minutes),
                    "type": day_type,
                })
                if shown_minutes:
                    work_days += 1
                    work_minutes += int(shown_minutes)
                continue

            if actual is not None and actual > 0:
                cells.append({
                    "mark": "Р",
                    "hours": format_hours_minutes(actual),
                    "type": day_type or "Робота",
                })
                work_days += 1
                work_minutes += int(actual)
                overtime_minutes += max(0, int(actual) - int(row["planned_minutes"] or 0))
                if d.weekday() >= 5:
                    weekend_minutes += int(actual)
            elif row["planned_minutes"] > 0:
                cells.append({"mark": "—", "hours": "", "type": day_type or "Робота"})
                missing_fact += 1
            else:
                cells.append({"mark": "", "hours": "", "type": day_type or "Вихідний"})

        rows.append({
            "employee_id": employee["id"],
            "personnel_no": employee["personnel_no"] or "",
            "name": core.employee_name(employee),
            "position": employee["position"] or "",
            "roles": employee["roles"] or "",
            "sex": "",
            "cells": cells,
            "work_days": work_days,
            "work_minutes": work_minutes,
            "overtime_minutes": overtime_minutes,
            "night_minutes": 0,
            "evening_minutes": 0,
            "weekend_minutes": weekend_minutes,
            "absence": absence,
            "missing_fact": missing_fact,
            "salary_rate": "",
        })

    con.close()
    return {
        "year": y,
        "month": m,
        "days": days,
        "company": company,
        "employees": rows,
    }


def _register_pdf_font(core, pdfmetrics, TTFont, name):
    path = next((p for p in core.report_font_candidates() if os.path.exists(p)), None)
    if path:
        try:
            if name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(name, path))
            return name
        except Exception:
            pass
    return "Helvetica"


def export_p5_pdf(
    core,
    year,
    month,
    out_path,
    *,
    active_only=True,
    form_date=None,
    department="",
    edrpou="",
):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    )

    data = collect_p5(core, year, month, active_only=active_only)
    form_date = form_date or date.today()
    font = _register_pdf_font(core, pdfmetrics, TTFont, "TaxoP5Font")
    styles = getSampleStyleSheet()
    normal = ParagraphStyle(
        "TaxoP5Normal", parent=styles["Normal"], fontName=font,
        fontSize=5.1, leading=5.8
    )
    center = ParagraphStyle(
        "TaxoP5Center", parent=normal, alignment=1, fontSize=5.0, leading=5.7
    )
    small = ParagraphStyle(
        "TaxoP5Small", parent=normal, fontSize=4.4, leading=5.0
    )
    title = ParagraphStyle(
        "TaxoP5Title", parent=center, fontSize=11.5, leading=13
    )

    doc = SimpleDocTemplate(
        str(out_path), pagesize=landscape(A4),
        leftMargin=9, rightMargin=9, topMargin=8, bottomMargin=8
    )
    story = []
    company_name = (data["company"]["name"] if data["company"] else "") or ""
    period_from = data["days"][0].strftime("%d.%m.%Y")
    period_to = data["days"][-1].strftime("%d.%m.%Y")

    head = Table(
        [
            [
                Paragraph(
                    f"<b>{escape(company_name)}</b><br/>"
                    f"{escape(department or '')}<br/>"
                    f"Ідентифікаційний код ЄДРПОУ: {escape(edrpou or '')}",
                    normal,
                ),
                Paragraph(
                    "<b>Типова форма № П-5</b><br/>"
                    "ЗАТВЕРДЖЕНО<br/>"
                    "Наказ Держкомстату України<br/>"
                    "05.12.2008 № 489",
                    normal,
                ),
            ],
            [
                Paragraph("<b>ТАБЕЛЬ ОБЛІКУ ВИКОРИСТАННЯ РОБОЧОГО ЧАСУ</b>", title),
                Paragraph(
                    f"Дата заповнення: <b>{form_date.strftime('%d.%m.%Y')}</b><br/>"
                    f"Звітний період: <b>{period_from}</b> - <b>{period_to}</b>",
                    normal,
                ),
            ],
        ],
        colWidths=[615, 205],
    )
    head.setStyle(TableStyle([
        ("FONTNAME", (0,0), (-1,-1), font),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("BOX", (1,0), (1,-1), 0.35, colors.black),
        ("GRID", (1,0), (1,-1), 0.25, colors.black),
        ("ALIGN", (1,0), (1,-1), "CENTER"),
        ("SPAN", (0,1), (0,1)),
        ("LEFTPADDING", (0,0), (-1,-1), 2),
        ("RIGHTPADDING", (0,0), (-1,-1), 2),
    ]))
    story.append(head)
    story.append(Spacer(1, 5))

    left_codes = P5_CODES[:15]
    right_codes = P5_CODES[15:]
    code_rows = [[
        Paragraph("<b>Умовні позначення</b>", center),
        Paragraph("<b>Код<br/>буквений</b>", center),
        Paragraph("<b>цифровий</b>", center),
        Paragraph("<b>Умовні позначення</b>", center),
        Paragraph("<b>Код<br/>буквений</b>", center),
        Paragraph("<b>цифровий</b>", center),
    ]]
    for idx in range(15):
        l = left_codes[idx]
        r = right_codes[idx] if idx < len(right_codes) else ("", "", "")
        code_rows.append([
            Paragraph(escape(l[0]), small), l[1], l[2],
            Paragraph(escape(r[0]), small), r[1], r[2],
        ])
    code_table = Table(
        code_rows,
        colWidths=[300, 38, 38, 300, 38, 38],
        repeatRows=1,
    )
    code_table.setStyle(TableStyle([
        ("FONTNAME", (0,0), (-1,-1), font),
        ("FONTSIZE", (0,0), (-1,-1), 4.8),
        ("GRID", (0,0), (-1,-1), 0.35, colors.black),
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#F2F2F2")),
        ("ALIGN", (1,0), (2,-1), "CENTER"),
        ("ALIGN", (4,0), (5,-1), "CENTER"),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("LEFTPADDING", (0,0), (-1,-1), 1.5),
        ("RIGHTPADDING", (0,0), (-1,-1), 1.5),
        ("TOPPADDING", (0,0), (-1,-1), 1.4),
        ("BOTTOMPADDING", (0,0), (-1,-1), 1.4),
    ]))
    story.append(code_table)

    story.append(PageBreak())

    day_labels = [f"{d.day:02d}" for d in data["days"]]
    summary_headers = ["Днів", "Год.", "НУ", "РН", "ВЧ", "РВ"]
    absence_headers = [label for _key, label in P5_ABSENCE_BUCKETS]
    headers = [
        "№", "Таб. №", "Стать", "ПІБ, посада",
        *day_labels, *summary_headers, *absence_headers, "Оклад / ставка"
    ]
    rows = [[Paragraph(f"<b>{escape(h)}</b>", center) for h in headers]]
    for idx, emp in enumerate(data["employees"], 1):
        day_cells = []
        for cell in emp["cells"]:
            value = cell["mark"]
            if cell["hours"]:
                value += f"<br/>{cell['hours']}"
            day_cells.append(Paragraph(value, center))
        absence_vals = []
        for key, _label in P5_ABSENCE_BUCKETS:
            stat = emp["absence"][key]
            absence_vals.append(
                f"{stat['days']}" + (
                    f"/{format_hours_minutes(stat['minutes'])}"
                    if stat["minutes"] else ""
                )
                if stat["days"] else ""
            )
        values = [
            idx,
            emp["personnel_no"],
            emp["sex"],
            Paragraph(
                f"<b>{escape(emp['name'])}</b><br/>{escape(emp['position'])}",
                small,
            ),
            *day_cells,
            emp["work_days"],
            format_hours_minutes(emp["work_minutes"]),
            format_hours_minutes(emp["overtime_minutes"]),
            format_hours_minutes(emp["night_minutes"]),
            format_hours_minutes(emp["evening_minutes"]),
            format_hours_minutes(emp["weekend_minutes"]),
            *absence_vals,
            emp["salary_rate"],
        ]
        rows.append(values)

    day_w = 9.2
    widths = [16, 43, 18, 91] + [day_w] * len(data["days"])
    widths += [23, 28, 22, 22, 22, 22]
    widths += [16] * len(P5_ABSENCE_BUCKETS)
    widths += [32]
    p5_table = Table(rows, colWidths=widths, repeatRows=1)
    style = [
        ("FONTNAME", (0,0), (-1,-1), font),
        ("FONTSIZE", (0,0), (-1,-1), 4.4),
        ("GRID", (0,0), (-1,-1), 0.28, colors.black),
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#EFEFEF")),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("ALIGN", (0,0), (2,-1), "CENTER"),
        ("ALIGN", (4,0), (-1,-1), "CENTER"),
        ("LEFTPADDING", (0,0), (-1,-1), 0.6),
        ("RIGHTPADDING", (0,0), (-1,-1), 0.6),
        ("TOPPADDING", (0,0), (-1,-1), 1.2),
        ("BOTTOMPADDING", (0,0), (-1,-1), 1.2),
    ]
    p5_table.setStyle(TableStyle(style))
    story.append(Paragraph(
        f"<b>ТАБЕЛЬ ОБЛІКУ ВИКОРИСТАННЯ РОБОЧОГО ЧАСУ - "
        f"{core.month_name_ua(data['month']).upper()} {data['year']}</b>",
        title,
    ))
    story.append(Spacer(1, 3))
    story.append(p5_table)
    story.append(Spacer(1, 4))
    story.append(Paragraph(
        "Позначка «—» у днях означає: робочий план є, але фактичний час у Taxo ще не внесено. "
        "Колонки причин неявок подані як дні/години; години беруться з перекритого робочого плану, якщо він був.",
        small,
    ))
    signer = (data["company"]["signer_name"] if data["company"] else "") or ""
    story.append(Spacer(1, 6))
    story.append(Paragraph(
        f'"{form_date.day:02d}" {core.month_name_ua(form_date.month).lower()} {form_date.year} р. '
        f"    Відповідальна особа: {escape(signer)} ____________________",
        normal,
    ))
    doc.build(story)
    return data


def export_p5_xlsx(
    core,
    year,
    month,
    out_path,
    *,
    active_only=True,
    form_date=None,
    department="",
    edrpou="",
):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    data = collect_p5(core, year, month, active_only=active_only)
    form_date = form_date or date.today()
    wb = Workbook()
    ws_codes = wb.active
    ws_codes.title = "Умовні позначення"
    ws = wb.create_sheet("Табель П-5")

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="EDEDED")

    ws_codes.merge_cells("A1:F1")
    ws_codes["A1"] = company_name = (
        (data["company"]["name"] if data["company"] else "") or ""
    )
    ws_codes["A1"].font = Font(size=12, bold=True)
    ws_codes["A1"].alignment = Alignment(horizontal="center")

    ws_codes.merge_cells("A2:D2")
    ws_codes["A2"] = department or ""
    ws_codes.merge_cells("A3:D3")
    ws_codes["A3"] = f"Ідентифікаційний код ЄДРПОУ: {edrpou or ''}"
    ws_codes["E1"] = "Типова форма № П-5"
    ws_codes["E2"] = "ЗАТВЕРДЖЕНО"
    ws_codes["E3"] = "Наказ Держкомстату України 05.12.2008 № 489"
    ws_codes.merge_cells("A5:F5")
    ws_codes["A5"] = "ТАБЕЛЬ ОБЛІКУ ВИКОРИСТАННЯ РОБОЧОГО ЧАСУ"
    ws_codes["A5"].font = Font(size=14, bold=True)
    ws_codes["A5"].alignment = Alignment(horizontal="center")
    ws_codes["A6"] = "Дата заповнення"
    ws_codes["B6"] = form_date.strftime("%d.%m.%Y")
    ws_codes["D6"] = "Звітний період"
    ws_codes["E6"] = data["days"][0].strftime("%d.%m.%Y")
    ws_codes["F6"] = data["days"][-1].strftime("%d.%m.%Y")

    headers = ["Умовні позначення", "Букв.", "Цифр.", "Умовні позначення", "Букв.", "Цифр."]
    for col, value in enumerate(headers, 1):
        cell = ws_codes.cell(8, col, value)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left, right = P5_CODES[:15], P5_CODES[15:]
    for i in range(15):
        l = left[i]
        r = right[i] if i < len(right) else ("", "", "")
        vals = [l[0], l[1], l[2], r[0], r[1], r[2]]
        for col, value in enumerate(vals, 1):
            c = ws_codes.cell(9 + i, col, value)
            c.border = border
            c.alignment = Alignment(
                horizontal="center" if col in (2,3,5,6) else "left",
                vertical="center", wrap_text=True
            )
        ws_codes.row_dimensions[9 + i].height = 34
    for col, width in enumerate([55, 9, 9, 55, 9, 9], 1):
        ws_codes.column_dimensions[get_column_letter(col)].width = width
    ws_codes.sheet_view.showGridLines = False
    ws_codes.page_setup.orientation = "landscape"
    ws_codes.page_setup.paperSize = ws_codes.PAPERSIZE_A4
    ws_codes.page_setup.fitToWidth = 1
    ws_codes.page_setup.fitToHeight = 1
    ws_codes.sheet_properties.pageSetUpPr.fitToPage = True
    ws_codes.print_area = "A1:F23"

    days = data["days"]
    summary_headers = ["Днів", "Год.", "НУ", "РН", "ВЧ", "РВ"]
    absence_headers = [label.replace("\n", " ") for _key, label in P5_ABSENCE_BUCKETS]
    table_headers = [
        "№ п/п", "Табельний номер", "Стать (ч/ж)", "ПІБ, посада",
        *[f"{d.day:02d}" for d in days],
        *summary_headers, *absence_headers, "Оклад, тарифна ставка, грн"
    ]

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(table_headers))
    ws.cell(1,1,company_name)
    ws.cell(1,1).font = Font(size=12, bold=True)
    ws.cell(1,1).alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(table_headers))
    ws.cell(2,1,f"ТАБЕЛЬ ОБЛІКУ ВИКОРИСТАННЯ РОБОЧОГО ЧАСУ - {core.month_name_ua(data['month']).upper()} {data['year']}")
    ws.cell(2,1).font = Font(size=14, bold=True)
    ws.cell(2,1).alignment = Alignment(horizontal="center")

    for col, value in enumerate(table_headers, 1):
        c = ws.cell(4, col, value)
        c.font = Font(size=8, bold=True)
        c.fill = header_fill
        c.border = border
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[4].height = 60

    row_no = 5
    for idx, emp in enumerate(data["employees"], 1):
        values = [idx, emp["personnel_no"], emp["sex"], f"{emp['name']}\n{emp['position']}"]
        for cell in emp["cells"]:
            value = cell["mark"]
            if cell["hours"]:
                value += f"\n{cell['hours']}"
            values.append(value)
        values += [
            emp["work_days"],
            format_hours_minutes(emp["work_minutes"]),
            format_hours_minutes(emp["overtime_minutes"]),
            format_hours_minutes(emp["night_minutes"]),
            format_hours_minutes(emp["evening_minutes"]),
            format_hours_minutes(emp["weekend_minutes"]),
        ]
        for key, _label in P5_ABSENCE_BUCKETS:
            stat = emp["absence"][key]
            values.append(
                (f"{stat['days']}/{format_hours_minutes(stat['minutes'])}"
                 if stat["days"] and stat["minutes"]
                 else str(stat["days"]) if stat["days"] else "")
            )
        values.append(emp["salary_rate"])

        for col, value in enumerate(values, 1):
            c = ws.cell(row_no, col, value)
            c.border = border
            c.alignment = Alignment(
                horizontal="left" if col == 4 else "center",
                vertical="center", wrap_text=True
            )
            c.font = Font(size=8)
        ws.row_dimensions[row_no].height = 28
        row_no += 1

    total_row = row_no
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=4)
    ws.cell(total_row,1,"Всього")
    ws.cell(total_row,1).font = Font(bold=True)
    ws.cell(total_row,1).alignment = Alignment(horizontal="right")
    total_work_days = sum(e["work_days"] for e in data["employees"])
    total_work_minutes = sum(e["work_minutes"] for e in data["employees"])
    first_summary_col = 5 + len(days)
    ws.cell(total_row, first_summary_col, total_work_days)
    ws.cell(total_row, first_summary_col + 1, format_hours_minutes(total_work_minutes))
    for col in range(1, len(table_headers) + 1):
        ws.cell(total_row, col).border = border

    widths = [6, 13, 8, 28] + [4.2] * len(days)
    widths += [7, 8, 7, 7, 7, 7] + [9] * len(P5_ABSENCE_BUCKETS) + [12]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.freeze_panes = "E5"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:4"
    ws.print_area = f"A1:{get_column_letter(len(table_headers))}{total_row}"
    ws.page_margins.left = 0.15
    ws.page_margins.right = 0.15
    ws.page_margins.top = 0.25
    ws.page_margins.bottom = 0.25

    wb.save(str(out_path))
    return data
